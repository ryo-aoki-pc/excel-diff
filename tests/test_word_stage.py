"""Exercise Word partial staging against real, isolated Git indexes."""

import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import shlex
import shutil
import subprocess
import sys
import tempfile
import unittest
from unittest import mock

from word_fixtures import (PNG, default_styles, hyperlink, image, make_docx,
                           numbering_xml, p, run, table, unpack, zip_parts)

SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "excel-stage.py"


def isolated_env(home):
    env = {key: value for key, value in os.environ.items() if not key.startswith("GIT_")}
    env.update(HOME=str(home), XDG_CONFIG_HOME=str(home / ".config"),
               GIT_CONFIG_NOSYSTEM="1", GIT_CONFIG_GLOBAL=os.devnull,
               GIT_TERMINAL_PROMPT="0", PYTHONUTF8="1")
    return env


def stage_document(main_type="docx"):
    body = p("架空の試験文書", style="1") + p("変更する本文です。") + p("削除候補の本文です。")
    body += table([["品目", "数量"], ["合成部品", "10"]]) + p("強調する本文です。")
    return make_docx(body, headers={"default": p("合成ヘッダー・初版")},
                     footers={"default": p("合成フッター・初版")}, main_type=main_type)


def replace_text(data, replacements):
    parts = unpack(data)
    for name, value in list(parts.items()):
        if name.endswith(".xml"):
            for before, after in replacements.items():
                value = value.replace((">" + before + "</w:t>").encode("utf-8"),
                                      (">" + after + "</w:t>").encode("utf-8"))
            parts[name] = value
    return zip_parts(parts)


@unittest.skipUnless(shutil.which("git"), "Git is required for partial staging tests")
class WordStageTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.scripts_path = str(SCRIPT.parent)
        sys.path.insert(0, cls.scripts_path)
        spec = importlib.util.spec_from_file_location("word_stage_integration", SCRIPT)
        cls.stage = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = cls.stage
        spec.loader.exec_module(cls.stage)
        import word_ooxml
        cls.word = word_ooxml

    @classmethod
    def tearDownClass(cls):
        sys.path.remove(cls.scripts_path)
        sys.modules.pop("word_stage_integration", None)

    def setUp(self):
        temporary = tempfile.TemporaryDirectory(prefix="Word 部分ステージ ")
        self.addCleanup(temporary.cleanup)
        self.directory = Path(temporary.name)
        self.repo = self.directory / "日本語 リポジトリ"
        self.repo.mkdir()
        home = self.directory / "home"
        home.mkdir()
        self.env = isolated_env(home)
        environment = mock.patch.dict(os.environ, self.env, clear=True)
        environment.start()
        self.addCleanup(environment.stop)
        self.git("init", "--quiet")
        self.git("config", "user.name", "Word staging test")
        self.git("config", "user.email", "word-stage@example.invalid")
        self.git("config", "core.autocrlf", "false")
        self.path = self.repo / "日本語 と空白.docx"
        self.path.write_bytes(stage_document())
        (self.repo / "notes.txt").write_bytes(b"original\n")
        self.git("add", "--", self.path.name, "notes.txt")
        self.git("commit", "--quiet", "-m", "Initial document")

    def git(self, *args, cwd=None, input=None, env=None, expected_returncode=0):
        result = subprocess.run(["git", "--no-pager", *args], cwd=cwd or self.repo,
                                env=env or self.env, input=input, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, expected_returncode, result.stderr.decode(errors="replace"))
        return result.stdout

    def edit(self, replacements, path=None):
        path = path or self.path
        path.write_bytes(replace_text(path.read_bytes(), replacements))

    def prepare(self, name="session"):
        directory = self.directory / name
        return directory, self.stage.prepare_session(self.repo, self.path.name, directory)

    def select(self, directory, state, predicate):
        """Stage complete replacements accepted by the requested predicate."""
        picker = directory / "picker"
        env = {key: value for key, value in self.env.items() if key != "GIT_INDEX_FILE"}
        selected = set()
        for part in state["parts"]:
            filename = part["filename"]
            before = self.git("show", "HEAD:" + filename, cwd=picker, env=env).decode("utf-8").splitlines()
            after = (picker / filename).read_text(encoding="utf-8").splitlines()
            result = []
            self.assertEqual(len(before), len(after))
            for old, new in zip(before, after):
                old_record, new_record = json.loads(old), json.loads(new)
                if old != new and predicate(part, old_record, new_record):
                    result.append(new)
                    selected.add((part["id"], new_record[0].split("/", 1)[0]))
                else:
                    result.append(old)
            (picker / filename).write_bytes(("\n".join(result) + "\n").encode("utf-8"))
            self.git("add", "--", filename, cwd=picker, env=env)
        return selected

    def select_text(self, directory, state, *texts):
        return self.select(directory, state, lambda part, old, new: new[3] in texts)

    def staged_bytes(self, path=None, cwd=None):
        return self.git("show", ":" + (path or self.path).name, cwd=cwd)

    def index_bytes(self):
        index = self.git("rev-parse", "--git-path", "index").decode().strip()
        return (self.repo / index).read_bytes()

    def rendered(self, data=None):
        return self.word.render_text(self.staged_bytes() if data is None else data)

    def assert_rejected(self, operation, pattern=None):
        before = self.index_bytes()
        with self.assertRaisesRegex(ValueError, pattern or "."):
            operation()
        self.assertEqual(self.index_bytes(), before)

    def test_body_header_footer_and_cell_selection_preserves_unselected_content(self):
        self.edit({"変更する本文です。": "変更した本文です。", "合成ヘッダー・初版": "合成ヘッダー・改訂",
                   "合成フッター・初版": "合成フッター・改訂", "10": "25", "合成部品": "改訂部品"})
        working = self.path.read_bytes()
        directory, state = self.prepare()
        selected = self.select_text(directory, state, "変更した本文です。", "合成ヘッダー・改訂", "合成フッター・改訂", "25")
        self.assertEqual(len(selected), 4)
        self.assertEqual(self.stage.apply_session(directory), 4)
        for text in ("変更した本文です。", "合成ヘッダー・改訂", "合成フッター・改訂", "25", "合成部品"):
            self.assertIn(text, self.rendered())
        self.assertNotIn("改訂部品", self.rendered())
        self.assertEqual(self.path.read_bytes(), working)

    def test_paragraph_insert_and_delete_are_selectable(self):
        parts = unpack(self.path.read_bytes())
        document = parts["word/document.xml"].decode("utf-8").replace(p("削除候補の本文です。"), "")
        parts["word/document.xml"] = document.replace("<w:sectPr>", p("末尾に追加する合成段落です。") + "<w:sectPr>")
        self.path.write_bytes(zip_parts(parts))
        directory, state = self.prepare()
        selected = self.select(directory, state, lambda part, old, new:
                               old[3] == "削除候補の本文です。" or new[3] == "末尾に追加する合成段落です。")
        self.assertEqual(len(selected), 2)
        self.assertEqual(self.stage.apply_session(directory), 2)
        self.assertNotIn("削除候補の本文です。", self.rendered())
        self.assertIn("末尾に追加する合成段落です。", self.rendered())

    def test_formatting_only_change_is_visible_and_selectable(self):
        parts = unpack(self.path.read_bytes())
        parts["word/document.xml"] = parts["word/document.xml"].replace(
            p("強調する本文です。").encode("utf-8"),
            p(runs=run("強調する本文です。", '<w:b/><w:color w:val="008800"/>')).encode("utf-8"))
        self.path.write_bytes(zip_parts(parts))
        directory, state = self.prepare()
        self.assertIn("太字", self.git("diff", "--", cwd=directory / "picker").decode("utf-8"))
        self.assertEqual(len(self.select_text(directory, state, "強調する本文です。")), 1)
        self.assertEqual(self.stage.apply_session(directory), 1)
        self.assertIn("太字", self.rendered())
        self.assertIn("色#008800", self.rendered())

    def test_selected_images_hyperlinks_styles_and_numbering_copy_dependencies(self):
        self.path.write_bytes(make_docx(p("依存関係を確認する本文です。") + p("未選択の本文です。")))
        self.git("add", "--", self.path.name)
        self.git("commit", "--quiet", "-m", "Dependency baseline")
        styles = default_styles().replace("</w:styles>", '<w:style w:type="paragraph" w:styleId="SyntheticExtra"><w:name w:val="Synthetic Extra"/><w:basedOn w:val="a"/></w:style></w:styles>')
        props = '<w:numPr><w:ilvl w:val="0"/><w:numId w:val="7"/></w:numPr>'
        self.path.write_bytes(make_docx(p(style="SyntheticExtra", props=props,
                                        runs=run("依存関係を確認する本文です。") + image() + hyperlink())
                                        + p("未選択の改訂本文です。"), styles=styles, numbering=numbering_xml(),
                                        rels=[("rImage", "image", "media/synthetic.png"),
                                              ("rLink", "hyperlink", "https://example.invalid/synthetic", "External")],
                                        parts={"word/media/synthetic.png": PNG}))
        directory, state = self.prepare()
        self.assertEqual(len(self.select(directory, state, lambda part, old, new: "依存関係を確認する" in new[3])), 1)
        self.assertEqual(self.stage.apply_session(directory), 1)
        parts = unpack(self.staged_bytes())
        self.assertIn(PNG, parts.values())
        self.assertIn(b"https://example.invalid/synthetic", parts["word/_rels/document.xml.rels"])
        self.assertIn(b"Synthetic Extra", parts["word/styles.xml"])
        self.assertIn("word/numbering.xml", parts)
        self.assertIn("未選択の本文です。", self.rendered())
        self.assertNotIn("未選択の改訂本文です。", self.rendered())

    def test_existing_document_stage_and_other_file_stage_survive_apply(self):
        self.edit({"変更する本文です。": "すでにステージした本文です。", "10": "20"})
        self.git("add", "--", self.path.name)
        self.edit({"すでにステージした本文です。": "まだ未選択の本文です。", "20": "30"})
        directory, state = self.prepare()
        self.select_text(directory, state, "30")
        (self.repo / "notes.txt").write_bytes(b"staged after session\n")
        self.git("add", "--", "notes.txt")
        self.assertEqual(self.stage.apply_session(directory), 1)
        self.assertIn("すでにステージした本文です。", self.rendered())
        self.assertNotIn("まだ未選択の本文です。", self.rendered())
        self.assertEqual(self.git("show", ":notes.txt"), b"staged after session\n")

    def test_session_contains_word_snapshots_parts_and_picker_identity(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        directory, state = self.prepare()
        self.assertEqual(state["format"], "word")
        self.assertEqual(state["target"], "index")
        self.assertFalse(state["applied"])
        self.assertEqual(state["differences"], 1)
        self.assertEqual(state["unsupported"], [])
        self.assertEqual((directory / "base.docx").read_bytes(), self.staged_bytes())
        self.assertEqual((directory / "work.docx").read_bytes(), self.path.read_bytes())
        self.assertFalse((directory / "base.xlsx").exists())
        self.assertEqual(self.git("config", "user.name", cwd=directory / "picker"), b"Word staging\n")
        self.assertEqual(self.git("config", "user.email", cwd=directory / "picker"), b"word-stage@example.invalid\n")
        self.assertEqual(self.git("log", "-1", "--format=%s", cwd=directory / "picker").decode("utf-8").strip(), "現在の Word インデックス")
        for number, part in enumerate(state["parts"], 1):
            self.assertTrue(part["filename"].startswith("{:03d}-".format(number)))
            self.assertTrue(part["filename"].endswith(".word.txt"))
            baseline = (directory / "baseline" / part["filename"]).read_bytes()
            self.assertEqual(hashlib.sha256(baseline).hexdigest(), part["baseline_sha256"])
            self.assertEqual(baseline, self.git("show", "HEAD:" + part["filename"], cwd=directory / "picker"))

    def test_japanese_picker_names_disable_quoting_only_in_picker(self):
        self.git("config", "core.quotePath", "true")
        self.edit({"変更する本文です。": "変更した本文です。"})
        directory, state = self.prepare()
        self.assertIn("本文", state["parts"][0]["filename"])
        self.assertEqual(self.git("config", "core.quotePath", cwd=directory / "picker"), b"false\n")
        self.assertEqual(self.git("config", "core.quotePath"), b"true\n")
        diff = self.git("diff", "--", cwd=directory / "picker")
        self.assertIn(state["parts"][0]["filename"].encode("utf-8"), diff)
        self.assertNotIn(b"\\345", diff)

    def test_apply_session_cli_reports_word_places(self):
        self.edit({"変更する本文です。": "CLI で変更する本文です。"})
        directory, state = self.prepare()
        self.select_text(directory, state, "CLI で変更する本文です。")
        result = subprocess.run([sys.executable, str(SCRIPT), "--apply-session", str(directory)],
                                cwd=self.repo, env=self.env, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8"))
        self.assertIn("1 か所の変更を Word のままステージしました。", result.stdout.decode("utf-8"))
        self.assertIn("CLI で変更する本文です。", self.rendered())

    def test_cli_without_selection_uses_word_selection_message(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        directory, _ = self.prepare()
        output, errors = io.StringIO(), io.StringIO()
        with mock.patch.object(self.stage.sys, "stdout", output), mock.patch.object(self.stage.sys, "stderr", errors):
            self.assertEqual(self.stage.main(["--apply-session", str(directory)]), 0)
        self.assertEqual(errors.getvalue(), "")
        self.assertIn("変更が選択されていません。Enter・Space", output.getvalue())
        self.assertNotIn("セル変更", output.getvalue())

    def test_linked_worktree_updates_its_own_index(self):
        worktree = self.directory / "別の 作業ツリー"
        self.git("worktree", "add", "--quiet", "-b", "word-partial-staging", str(worktree))
        before = self.index_bytes()
        work_path = worktree / self.path.name
        self.edit({"変更する本文です。": "作業ツリーで変更する本文です。"}, path=work_path)
        directory = self.directory / "linked-session"
        state = self.stage.prepare_session(worktree, work_path.name, directory)
        self.assertEqual(Path(state["repo"]), worktree)
        self.assertNotEqual(Path(state["index"]), self.repo / ".git" / "index")
        self.select_text(directory, state, "作業ツリーで変更する本文です。")
        self.assertEqual(self.stage.apply_session(directory), 1)
        self.assertIn("作業ツリーで変更する本文です。", self.rendered(self.staged_bytes(work_path, cwd=worktree)))
        self.assertEqual(self.index_bytes(), before)

    def test_alternate_index_is_respected_and_isolated_from_picker(self):
        before = self.index_bytes()
        alternate = self.directory / "別の インデックス"
        alternate.write_bytes(before)
        self.edit({"変更する本文です。": "別インデックスの本文です。"})
        with mock.patch.dict(self.env, {"GIT_INDEX_FILE": str(alternate)}), mock.patch.dict(os.environ, {"GIT_INDEX_FILE": str(alternate)}):
            directory, state = self.prepare()
            self.assertEqual(Path(state["index"]), alternate)
            self.select_text(directory, state, "別インデックスの本文です。")
            (self.repo / "notes.txt").write_bytes(b"alternate index change\n")
            self.git("add", "--", "notes.txt")
            self.assertEqual(self.stage.apply_session(directory), 1)
            self.assertIn("別インデックスの本文です。", self.rendered())
            self.assertEqual(self.git("show", ":notes.txt"), b"alternate index change\n")
        self.assertEqual(self.index_bytes(), before)

    def test_literal_filename_pathspec_characters_and_leading_dash(self):
        filenames = ("-日本語 文書.docx", "literal[1].docx")
        if os.name != "nt":
            filenames += ("literal*?.docx", ":(glob)document.docx")
        for number, name in enumerate(filenames):
            with self.subTest(name=name):
                path = self.repo / name
                path.write_bytes(stage_document())
                env = dict(self.env, GIT_LITERAL_PATHSPECS="1")
                self.git("add", "--", name, env=env)
                self.git("commit", "--quiet", "-m", "Literal filename fixture")
                self.edit({"変更する本文です。": "名前を確認する本文です。"}, path=path)
                directory = self.directory / ("literal-" + str(number))
                state = self.stage.prepare_session(self.repo, name, directory)
                self.select_text(directory, state, "名前を確認する本文です。")
                self.assertEqual(self.stage.apply_session(directory), 1)
                self.assertIn("名前を確認する本文です。", self.rendered(self.staged_bytes(path)))

    def test_cancel_and_apply_without_selection_preserve_index_exactly(self):
        self.edit({"変更する本文です。": "選ばない本文です。"})
        before = self.index_bytes()
        directory, _ = self.prepare()
        self.assertEqual(self.index_bytes(), before)
        self.assertEqual(self.stage.apply_session(directory), 0)
        self.assertEqual(self.index_bytes(), before)

    def test_target_index_change_rejects_newer_stage(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        directory, state = self.prepare()
        self.select_text(directory, state, "変更した本文です。")
        self.git("add", "--", self.path.name)
        self.assert_rejected(lambda: self.stage.apply_session(directory), "対象 Word のインデックスが変更")

    def test_working_file_change_and_deletion_reject(self):
        original = self.path.read_bytes()
        for deleted in (False, True):
            with self.subTest(deleted=deleted):
                self.path.write_bytes(replace_text(original, {"変更する本文です。": "変更した本文です。"}))
                directory, state = self.prepare("working-" + str(deleted))
                self.select_text(directory, state, "変更した本文です。")
                if deleted:
                    self.path.unlink()
                else:
                    self.edit({"合成部品": "新しい部品"})
                self.assert_rejected(lambda: self.stage.apply_session(directory), "作業ファイルが選択開始後に変更")

    def test_existing_index_lock_is_preserved(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        directory, state = self.prepare()
        self.select_text(directory, state, "変更した本文です。")
        lock = self.repo / ".git" / "index.lock"
        lock.write_bytes(b"another Git operation")
        self.assert_rejected(lambda: self.stage.apply_session(directory), "別の操作が使用中")
        self.assertEqual(lock.read_bytes(), b"another Git operation")

    def test_already_applied_session_is_rejected(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        directory, state = self.prepare()
        self.select_text(directory, state, "変更した本文です。")
        self.assertEqual(self.stage.apply_session(directory), 1)
        self.assert_rejected(lambda: self.stage.apply_session(directory), "反映済み")

    def test_missing_duplicate_reordered_readdressed_and_invented_records_are_rejected(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        for malformed in ("missing", "duplicate", "reordered", "readdressed", "invented"):
            with self.subTest(malformed=malformed):
                directory, state = self.prepare(malformed)
                filename = state["parts"][0]["filename"]
                lines = self.git("show", "HEAD:" + filename, cwd=directory / "picker").decode("utf-8").splitlines()
                if malformed == "missing":
                    lines.pop(0)
                elif malformed == "duplicate":
                    lines.append(lines[0])
                elif malformed == "reordered":
                    lines[0], lines[1] = lines[1], lines[0]
                else:
                    record = json.loads(lines[0])
                    record[0 if malformed == "readdressed" else 3] = "999" if malformed == "readdressed" else "どちらにもない値"
                    lines[0] = json.dumps(record, ensure_ascii=False, separators=(",", ":"))
                (directory / "picker" / filename).write_bytes(("\n".join(lines) + "\n").encode("utf-8"))
                self.git("add", "--", filename, cwd=directory / "picker")
                self.assert_rejected(lambda: self.stage.apply_session(directory))

    def test_picker_file_addition_deletion_and_conflict_are_rejected(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        for malformed in ("added", "deleted", "conflict"):
            with self.subTest(malformed=malformed):
                directory, state = self.prepare(malformed)
                picker = directory / "picker"
                filename = state["parts"][0]["filename"]
                if malformed == "added":
                    (picker / "extra.word.txt").write_bytes(b"[]\n")
                    self.git("add", "--", "extra.word.txt", cwd=picker)
                elif malformed == "deleted":
                    self.git("rm", "--force", "--", filename, cwd=picker)
                else:
                    oid = self.git("rev-parse", "HEAD:" + filename, cwd=picker).decode().strip()
                    self.git("update-index", "--force-remove", "--", filename, cwd=picker)
                    rows = "".join("100644 {} {}\t{}\n".format(oid, stage, filename) for stage in (1, 2, 3))
                    self.git("update-index", "--index-info", input=rows.encode("utf-8"), cwd=picker)
                self.assert_rejected(lambda: self.stage.apply_session(directory), "選択画面のファイルが削除・追加・競合")

    def test_changed_snapshots_and_saved_projection_are_rejected(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        for malformed in ("base.docx", "work.docx", "projection"):
            with self.subTest(malformed=malformed):
                directory, state = self.prepare(malformed)
                self.select_text(directory, state, "変更した本文です。")
                if malformed == "projection":
                    (directory / "baseline" / state["parts"][0]["filename"]).write_bytes(b"tampered\n")
                else:
                    self.edit({"合成部品": "別の部品"}, path=directory / malformed)
                self.assert_rejected(lambda: self.stage.apply_session(directory), "画面を閉じて選び直して")

    def test_rehashed_projection_must_match_fresh_word_comparison(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        directory, state = self.prepare()
        part = state["parts"][0]
        changed = (directory / "picker" / part["filename"]).read_bytes()
        (directory / "baseline" / part["filename"]).write_bytes(changed)
        part["baseline_sha256"] = hashlib.sha256(changed).hexdigest()
        self.stage.write_json(directory / "session.json", state)
        self.assert_rejected(lambda: self.stage.apply_session(directory), "射影が Word 文書と一致しません")

    def test_legacy_word_session_requires_reopening(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        for missing in ("base_sha256", "baseline_sha256"):
            with self.subTest(missing=missing):
                directory, state = self.prepare(missing)
                if missing == "baseline_sha256":
                    state["parts"][0].pop(missing)
                else:
                    state.pop(missing)
                self.stage.write_json(directory / "session.json", state)
                self.assert_rejected(lambda: self.stage.apply_session(directory), "画面を閉じて選び直して")

    def test_all_word_extensions_accept_partial_staging_and_preserve_vba(self):
        for extension in ("docx", "docm", "dotx", "dotm"):
            with self.subTest(extension=extension):
                path = self.repo / ("形式." + extension)
                original = stage_document(extension)
                path.write_bytes(original)
                self.git("add", "--", path.name)
                self.git("commit", "--quiet", "-m", "Word type fixture")
                self.edit({"変更する本文です。": "形式を確認する本文です。"}, path=path)
                directory = self.directory / extension
                state = self.stage.prepare_session(self.repo, path.name, directory)
                self.select_text(directory, state, "形式を確認する本文です。")
                self.assertEqual(self.stage.apply_session(directory), 1)
                data = self.staged_bytes(path)
                self.assertIn("形式を確認する本文です。", self.rendered(data))
                if extension in ("docm", "dotm"):
                    self.assertEqual(unpack(data)["word/vbaProject.bin"], unpack(original)["word/vbaProject.bin"])

    def test_new_and_unmerged_word_documents_require_whole_file_staging(self):
        new = self.repo / "新規.docx"
        new.write_bytes(stage_document())
        for staged in (False, True):
            with self.subTest(staged=staged):
                if staged:
                    self.git("add", "--", new.name)
                self.assert_rejected(lambda: self.stage.prepare_session(self.repo, new.name, self.directory / ("new-" + str(staged))), "Word")
        oid = self.git("rev-parse", ":" + self.path.name).decode().strip()
        self.git("update-index", "--force-remove", "--", self.path.name)
        rows = "".join("100644 {} {}\t{}\n".format(oid, stage, self.path.name) for stage in (1, 2, 3))
        self.git("update-index", "--index-info", input=rows.encode("utf-8"))
        self.assert_rejected(self.prepare, "追跡済みで競合のない Word")

    def test_outside_deleted_corrupt_and_corrupt_index_documents_are_rejected(self):
        outside = self.directory / "外部.docx"
        outside.write_bytes(stage_document())
        self.assert_rejected(lambda: self.stage.prepare_session(self.repo, str(outside), self.directory / "outside"), "リポジトリ内の Word")
        self.path.unlink()
        self.assert_rejected(self.prepare)
        self.path.write_bytes(b"PK\x03\x04broken document")
        self.assert_rejected(self.prepare)
        self.path.write_bytes(b"not a Word package")
        self.git("add", "--", self.path.name)
        self.path.write_bytes(stage_document())
        self.assert_rejected(self.prepare)

    def test_unicode_record_separators_are_escaped_and_selectable(self):
        for number, separator in enumerate(("\x85", "\u2028", "\u2029")):
            with self.subTest(separator=repr(separator)):
                self.path.write_bytes(replace_text(stage_document(), {"変更する本文です。": "前" + separator + "後"}))
                directory, state = self.prepare("separator-" + str(number))
                raw = (directory / "picker" / state["parts"][0]["filename"]).read_bytes()
                self.assertNotIn(separator.encode("utf-8"), raw)
                self.select_text(directory, state, "前" + separator + "後")
                self.assertEqual(self.stage.apply_session(directory), 1)

    def unsupported_document(self, count=1, updated=False):
        body = p("取り込める本文です。")
        for number in range(count):
            text = "コメント付きの本文 {}・".format(number) + ("改訂" if updated else "初版")
            body += p(runs=run(text) + '<w:r><w:commentReference w:id="8"/></w:r>')
        comments = '<w:comment w:id="8" w:author="合成試験者">' + p("合成コメント") + '</w:comment>'
        return make_docx(body, comments=comments)

    def prepare_unsupported(self, count=1, mixed=False):
        self.path.write_bytes(self.unsupported_document(count=count))
        self.git("add", "--", self.path.name)
        self.git("commit", "--quiet", "-m", "Unsupported baseline")
        data = self.unsupported_document(count=count, updated=True)
        if mixed:
            data = replace_text(data, {"取り込める本文です。": "取り込める改訂本文です。"})
        self.path.write_bytes(data)

    def test_unsupported_changes_are_shown_raw_and_refused_with_location_and_reason(self):
        self.prepare_unsupported()
        directory, state = self.prepare()
        self.assertEqual(len(state["unsupported"]), 1)
        diff = self.git("diff", "--", cwd=directory / "picker").decode("utf-8")
        self.assertIn("改訂", diff)
        self.assertEqual(len(self.select(directory, state, lambda part, old, new: "改訂" in new[3])), 1)
        self.assert_rejected(lambda: self.stage.apply_session(directory), "本文 .*コメント")

    def test_supported_selection_preserves_unselected_unsupported_change(self):
        self.prepare_unsupported(mixed=True)
        directory, state = self.prepare()
        self.select_text(directory, state, "取り込める改訂本文です。")
        self.assertEqual(self.stage.apply_session(directory), 1)
        self.assertIn("取り込める改訂本文です。", self.rendered())
        self.assertIn("コメント付きの本文 0・初版", self.rendered())
        self.assertNotIn("コメント付きの本文 0・改訂", self.rendered())

    def test_main_prints_first_twenty_unsupported_locations_before_picker(self):
        self.prepare_unsupported(count=22)
        before = self.index_bytes()
        output, errors = io.StringIO(), io.StringIO()
        observed = []
        def pick(directory):
            observed.append(output.getvalue())
            return 0
        with mock.patch.object(self.stage, "check_tools"), mock.patch.object(self.stage.Path, "cwd", return_value=self.repo), mock.patch.object(self.stage, "run_picker", side_effect=pick), mock.patch.object(self.stage.sys, "stdout", output), mock.patch.object(self.stage.sys, "stderr", errors):
            self.assertEqual(self.stage.main([self.path.name]), 0)
        self.assertEqual(len(observed), 1)
        self.assertIn("Word を読み込んでいます。大きい文書では選択画面", observed[0])
        self.assertIn("部分ステージできない違い（22 か所）:", observed[0])
        self.assertIn("本文 #", observed[0])
        self.assertIn("、ほか 2 か所", observed[0])
        self.assertEqual(observed[0].count("（コメント"), 20)
        self.assertIn("選択を取り消しました。Word のインデックスは変更していません。", output.getvalue())
        self.assertEqual(errors.getvalue(), "")
        self.assertEqual(self.index_bytes(), before)

    def test_main_success_uses_word_texts_and_cleans_session(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        output, errors = io.StringIO(), io.StringIO()
        sessions = []
        def pick(directory):
            sessions.append(directory)
            state = json.loads((directory / "session.json").read_text(encoding="utf-8"))
            self.select_text(directory, state, "変更した本文です。")
            self.assertEqual(self.stage.apply_session(directory), 1)
            return 0
        with mock.patch.object(self.stage, "check_tools"), mock.patch.object(self.stage.Path, "cwd", return_value=self.repo), mock.patch.object(self.stage, "run_picker", side_effect=pick), mock.patch.object(self.stage.sys, "stdout", output), mock.patch.object(self.stage.sys, "stderr", errors):
            self.assertEqual(self.stage.main([self.path.name]), 0)
        self.assertIn("Word の部分ステージを反映しました。", output.getvalue())
        self.assertEqual(errors.getvalue(), "")
        self.assertFalse(sessions[0].exists())

    def test_word_picker_config_has_wrapped_renderer_first_and_word_action(self):
        self.edit({"変更する本文です。": "変更した本文です。"})
        directory, _ = self.prepare()
        config = json.loads(self.stage.picker_config(directory).read_text(encoding="utf-8"))
        self.assertEqual(config["git"]["diffRenderers"], [
            {"command": "delta --no-gitconfig --dark --paging=never --syntax-theme=none --tabs=4 --side-by-side --wrap-max-lines=unlimited", "name": "delta side-by-side wrapped"},
            {"command": "delta --no-gitconfig --dark --paging=never --syntax-theme=none --tabs=4 --side-by-side", "name": "delta side-by-side"},
            {"type": "rawGit", "name": "default"},
        ])
        self.assertEqual(config["customCommands"][0]["description"], "選択した変更を Word へステージ (反映後 q で戻る)")
        self.assertEqual(config["customCommands"][0]["outputTitle"], "Word の部分ステージ")
        self.assertTrue(config["gui"]["useHunkModeInDiffView"])

    def test_excel_picker_config_equals_pre_word_literal(self):
        directory = self.directory / "excel-config"
        directory.mkdir()
        for merge in (False, True):
            with self.subTest(merge=merge):
                parts = [sys.executable, str(SCRIPT.resolve()), "--apply-merge-session" if merge else "--apply-session", str(directory)]
                command = subprocess.list2cmdline(parts) if os.name == "nt" else shlex.join(parts)
                expected = {
                    "gui": {"useHunkModeInDiffView": True, "language": "ja", "mouseEvents": False,
                            "showRandomTip": False, "showFileTree": False},
                    "git": {"autoFetch": False, "diffRenderers": [
                        {"command": "delta --no-gitconfig --dark --paging=never --syntax-theme=none --tabs=4 --side-by-side",
                         "name": "delta side-by-side"}, {"type": "rawGit", "name": "default"}]},
                    "update": {"method": "never"}, "disableStartupPopups": True,
                    "customCommands": [{"key": "X", "context": "global", "command": command,
                                        "description": ("選択セルを取り込み先へ書き込む (反映後 q で戻る)" if merge else
                                                        "選択セルを Excel へステージ (反映後 q で戻る)"),
                                        "output": "popup", "outputTitle": "Excel の取り込み" if merge else "Excel の部分ステージ"}],
                }
                self.assertEqual(json.loads(self.stage.picker_config(directory, merge=merge).read_text(encoding="utf-8")), expected)

    def test_equal_reference_swap_notice_follows_successful_index_write(self):
        def document(updated=False):
            name, ident = ("_TocStageNew", "8") if updated else ("_TocStageOld", "1")
            heading = p(runs='<w:bookmarkStart w:id="' + ident + '" w:name="' + name + '"/>'
                        + run("合成の見出し") + '<w:bookmarkEnd w:id="' + ident + '"/>')
            return make_docx(p(runs=hyperlink("更新後の参照" if updated else "更新前の参照", anchor=name))
                             + heading + p("末尾"))
        self.path.write_bytes(document())
        self.git("add", "--", self.path.name)
        self.git("commit", "--quiet", "-m", "Hidden reference baseline")
        self.path.write_bytes(document(True))
        directory, state = self.prepare()
        self.assertEqual(len(self.select_text(directory, state, "更新後の参照")), 1)
        output = io.StringIO()
        index = Path(state["index"])
        original_replace = self.stage.os.replace
        def replace(source, target):
            if Path(target) == index:
                raise PermissionError("index write blocked")
            return original_replace(source, target)
        before = self.index_bytes()
        with mock.patch.object(self.stage.os, "replace", side_effect=replace), mock.patch.object(
            self.stage.sys, "stdout", output
        ):
            with self.assertRaisesRegex(PermissionError, "index write blocked"):
                self.stage.apply_session(directory)
        self.assertEqual(self.index_bytes(), before)
        self.assertFalse(Path(str(index) + ".lock").exists())
        self.assertFalse(json.loads((directory / "session.json").read_text(encoding="utf-8"))["applied"])
        self.assertNotIn("目次などの参照先", output.getvalue())
        with mock.patch.object(self.stage.sys, "stdout", output):
            self.assertEqual(self.stage.apply_session(directory), 1)
        self.assertIn("目次などの参照先に合わせて、表示が同じ 1 か所を取り込み元のものにしました。", output.getvalue())
        self.assertIn(b"_TocStageNew", unpack(self.staged_bytes())["word/document.xml"])

    def test_copied_excel_helper_runs_without_word_modules_and_reports_word_import_error(self):
        from openpyxl import Workbook, load_workbook
        helpers = self.directory / "Word modules なし"
        helpers.mkdir()
        helper = helpers / SCRIPT.name
        shutil.copyfile(SCRIPT, helper)
        shutil.copyfile(SCRIPT.with_name("excel_stage_ooxml.py"), helpers / "excel_stage_ooxml.py")
        excel = self.repo / "単独 Excel.xlsx"
        book = Workbook()
        book.active["A1"] = "元の値"
        book.save(excel)
        self.git("add", "--", excel.name)
        self.git("commit", "--quiet", "-m", "Copied Excel helper baseline")
        book.active["A1"] = "変更後の値"
        book.save(excel)
        book.close()
        code = r'''
import importlib.util
from pathlib import Path
import sys
helper, repo, directory = map(Path, sys.argv[1:4])
sys.path.insert(0, str(helper.parent))
spec = importlib.util.spec_from_file_location("copied_stage", helper)
stage = importlib.util.module_from_spec(spec)
sys.modules[spec.name] = stage
spec.loader.exec_module(stage)
assert "word_ooxml" not in sys.modules
assert "word_stage_ooxml" not in sys.modules
state = stage.prepare_session(repo, sys.argv[4], directory)
assert "format" not in state
stage.git(directory / "picker", "add", "--all", env=stage.isolated_env())
assert stage.apply_session(directory) == (1, 0)  # values, formats
assert "word_ooxml" not in sys.modules
assert "word_stage_ooxml" not in sys.modules
try:
    stage._word_module()
except stage.StageError as error:
    assert str(error).startswith("Word 用のモジュールを読み込めません。setup-excel-diff.py を再実行してください: ")
    print(error)
else:
    raise AssertionError("Missing Word modules were accepted")
'''
        result = subprocess.run([sys.executable, "-I", "-X", "utf8", "-B", "-c", code,
                                 str(helper), str(self.repo), str(self.directory / "copied-session"), excel.name],
                                cwd=helpers, env=self.env, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr.decode("utf-8"))
        self.assertIn("Word 用のモジュールを読み込めません。", result.stdout.decode("utf-8"))
        staged = load_workbook(io.BytesIO(self.staged_bytes(excel)))
        self.addCleanup(staged.close)
        self.assertEqual(staged.active["A1"].value, "変更後の値")


if __name__ == "__main__":
    unittest.main()
