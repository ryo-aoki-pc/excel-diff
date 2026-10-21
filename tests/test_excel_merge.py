"""Exercise two-workbook imports without opening WinMerge or lazygit."""

import contextlib
import hashlib
import importlib.util
import io
import json
import os
from pathlib import Path
import shutil
import shlex
import stat
import subprocess
import sys
import tempfile
import types
import unittest
from unittest import mock
import zipfile
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts" / "excel-stage.py"
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"


def save_book(path):
    book = Workbook()
    sheet = book.active
    sheet.title = "売上"
    for coordinate, value in {"A1": "元の値", "B1": 10, "C1": True,
                              "D1": "=B1*2", "E1": "削除する値"}.items():
        sheet[coordinate] = value
    sheet["A1"].font = Font(bold=True, color="FF0000")
    sheet["A1"].fill = PatternFill("solid", fgColor="FFFF00")
    book.create_sheet("備考")["B4"] = "別シート"
    book.save(path)
    book.close()


def replace_parts(data, replacements):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(output, "w") as result:
        for info in original.infolist():
            result.writestr(info, replacements.pop(info.filename, original.read(info.filename)))
        for name, value in replacements.items():
            result.writestr(name, value)
    return output.getvalue()


def rich_cell(path, coordinate="A1"):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    cell = root.find(".//{" + NS + "}c[@r='" + coordinate + "']")
    for child in list(cell):
        cell.remove(child)
    cell.set("t", "inlineStr")
    inline = ET.SubElement(cell, "{" + NS + "}is")
    run = ET.SubElement(inline, "{" + NS + "}r")
    ET.SubElement(run, "{" + NS + "}t").text = "リッチテキスト"
    path.write_bytes(replace_parts(path.read_bytes(), {
        "xl/worksheets/sheet1.xml": ET.tostring(root, encoding="utf-8")
    }))


@unittest.skipUnless(shutil.which("git"), "Git is required for merge picker tests")
class ExcelMergeTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.scripts_path = str(SCRIPT.parent)
        sys.path.insert(0, cls.scripts_path)
        spec = importlib.util.spec_from_file_location("excel_merge_integration", SCRIPT)
        cls.stage = importlib.util.module_from_spec(spec)
        sys.modules[spec.name] = cls.stage
        spec.loader.exec_module(cls.stage)

    @classmethod
    def tearDownClass(cls):
        sys.path.remove(cls.scripts_path)
        sys.modules.pop("excel_merge_integration", None)

    def setUp(self):
        self.directory = Path(tempfile.mkdtemp(prefix="excel-merge-test-"))
        self.addCleanup(self.stage.cleanup_session, self.directory, retries=0, delay=0)
        self.target = self.directory / "取り込み先 と空白.xlsx"
        self.source = self.directory / "取り込み元 と空白.xlsx"
        save_book(self.target)
        self.source.write_bytes(self.target.read_bytes())
        self.edit(self.source, {("売上", "A1"): "新しい値", ("売上", "B1"): 25,
                                ("売上", "E1"): None, ("売上", "G1"): "追加",
                                ("備考", "B4"): "別シートの変更"})
        self.env = {key: value for key, value in os.environ.items() if not key.startswith("GIT_")}
        self.env.update(HOME=str(self.directory), XDG_CONFIG_HOME=str(self.directory / ".config"),
                        GIT_CONFIG_GLOBAL=os.devnull, GIT_CONFIG_NOSYSTEM="1", GIT_TERMINAL_PROMPT="0")
        environment = mock.patch.dict(os.environ, self.env, clear=True)
        environment.start()
        self.addCleanup(environment.stop)
        self.number = 0

    def edit(self, path, changes):
        book = load_workbook(path)
        for (sheet, coordinate), value in changes.items():
            book[sheet][coordinate] = value
        book.save(path)
        book.close()

    def git(self, *args, cwd=None, env=None, input=None):
        result = subprocess.run(["git", "--no-pager", *args], cwd=cwd or self.directory,
                                env=env or self.env, input=input, capture_output=True, timeout=30)
        self.assertEqual(result.returncode, 0, result.stderr.decode(errors="replace"))
        return result.stdout

    def prepare(self, tool="lazygit"):
        self.number += 1
        directory = self.directory / ("session-" + str(self.number))
        state = self.stage.prepare_merge_session(self.target, self.source, directory, tool=tool)
        return directory, state

    def select(self, directory, state, selected=(), formats=()):
        # Match test_excel_stage.py: stage complete record replacements only.
        picker = directory / "picker"
        for sheet in state["sheets"]:
            for key, selection in (("filename", selected), ("formats_filename", formats)):
                filename = sheet[key]
                baseline = self.git("show", "HEAD:" + filename, cwd=picker).decode("utf-8")
                working = (picker / filename).read_text(encoding="utf-8")
                # Blank lines separate format records and are never replaced.
                replacements = {json.loads(line)[0]: line for line in working.splitlines()
                                if line and (sheet["name"], json.loads(line)[0]) in selection}
                result = [replacements.get(json.loads(line)[0], line) if line else line for line in baseline.splitlines()]
                (picker / filename).write_text("\n".join(result) + ("\n" if result else ""), encoding="utf-8")
                self.git("add", "--", filename, cwd=picker)

    def win_paths(self, directory, state):
        return tuple(directory / state["winmerge"][key] for key in ("left", "right", "baseline"))

    def copied_lines(self, directory, state, selected=None, formats=()):
        left, right, _ = self.win_paths(directory, state)
        baseline = left.read_text(encoding="utf-8-sig").splitlines()
        changed = right.read_text(encoding="utf-8-sig").splitlines()
        sheet_name, format_section = None, False
        for index, line in enumerate(baseline):
            if line.startswith("# シート "):
                sheet_name = json.loads(line.split(": ", 1)[1])
                format_section = False
            elif line.startswith("# 書式 シート "):
                format_section = True
            elif line.startswith("["):
                selection = formats if format_section else selected
                if selection is None or (sheet_name, json.loads(line)[0]) in selection:
                    baseline[index] = changed[index]
        return baseline

    def run_win(self, directory, state, edit=None):
        left, _, _ = self.win_paths(directory, state)

        def run(command, *args, **kwargs):
            if edit is not None:
                data = edit(left.read_bytes())
                if data is not None:
                    left.write_bytes(data)
            return subprocess.CompletedProcess(command, 0)

        with mock.patch.object(self.stage, "find_winmerge", return_value="WinMergeU.exe"), mock.patch.object(
            self.stage.subprocess, "run", side_effect=run
        ) as launched, mock.patch.object(self.stage.sys, "stdin", io.StringIO()):
            selected = self.stage.run_winmerge(directory)
        return selected, launched

    def apply_win(self, selected=None, encoding="utf-8-sig", separator="\n", final="\n"):
        directory, state = self.prepare("winmerge")
        lines = self.copied_lines(directory, state, selected)
        chosen, launched = self.run_win(directory, state, lambda _: (separator.join(lines) + final).encode(encoding))
        self.assertEqual(len(chosen.values), len(selected) if selected is not None else 5)
        self.assertEqual(chosen.formats, set())
        self.assertEqual(self.stage.apply_merge_session(directory), (len(chosen.values), 0))
        return directory, state, launched

    def assert_unchanged_failure(self, operation, pattern=None):
        before = self.target.read_bytes() if self.target.is_file() else None
        with self.assertRaisesRegex((ValueError, OSError), pattern or "."):
            operation()
        if before is not None:
            self.assertEqual(self.target.read_bytes(), before)

    def main(self, arguments, picker=None, winmerge=None):
        output, errors = io.StringIO(), io.StringIO()
        with mock.patch.object(self.stage, "choose_tool", return_value="winmerge" if winmerge else "lazygit"), mock.patch.object(
            self.stage, "check_tools"
        ) as tools, mock.patch.object(self.stage, "run_picker", side_effect=picker) as lazy, mock.patch.object(
            self.stage, "run_winmerge", side_effect=winmerge
        ) as win, mock.patch.object(self.stage.sys, "stdout", output), mock.patch.object(
            self.stage.sys, "stderr", errors
        ), mock.patch.object(self.stage.sys, "stdin", io.StringIO()):
            status = self.stage.main(["--merge", *arguments])
        return status, output.getvalue(), errors.getvalue(), tools, lazy, win

    def test_prepare_preserves_both_inputs_and_builds_picker_projections(self):
        snapshots = [(path.read_bytes(), path.stat().st_mtime_ns) for path in (self.target, self.source)]
        directory, state = self.prepare()
        self.assertEqual(state["target"], "file")
        self.assertEqual(state["path"], str(self.target.resolve()))
        self.assertEqual(state["source"], str(self.source.resolve()))
        self.assertFalse(state["applied"])
        self.assertEqual(state["differences"], 5)
        self.assertEqual(state["format_differences"], 0)
        self.assertEqual(state["selectable"], 5)
        self.assertEqual(state["unsupported"], [])
        self.assertEqual((directory / "base.xlsx").read_bytes(), snapshots[0][0])
        self.assertEqual((directory / "work.xlsx").read_bytes(), snapshots[1][0])
        base, source = self.stage.Workbook(snapshots[0][0]), self.stage.Workbook(snapshots[1][0])
        format_base, format_source, _ = base.compare_formats(source).projection()
        for sheet in state["sheets"]:
            self.assertEqual(self.git("show", "HEAD:" + sheet["formats_filename"], cwd=directory / "picker").decode(),
                             format_base[sheet["id"]])
            self.assertEqual((directory / "picker" / sheet["formats_filename"]).read_text(encoding="utf-8"),
                             format_source[sheet["id"]])
            self.assertEqual(self.git("show", "HEAD:" + sheet["filename"], cwd=directory / "picker").decode(),
                             base.projection(source)[sheet["id"]])
            self.assertEqual((directory / "picker" / sheet["filename"]).read_text(encoding="utf-8"),
                             source.projection(base)[sheet["id"]])
        self.assertEqual(self.git("log", "-1", "--format=%s", cwd=directory / "picker").decode().strip(), "取り込み先の Excel")
        self.assertEqual([(path.read_bytes(), path.stat().st_mtime_ns) for path in (self.target, self.source)], snapshots)

    def test_formatting_only_difference_opens_picker_and_imports(self):
        original = self.target.read_bytes()
        original_run = subprocess.run
        for tool in ("lazygit", "winmerge"):
            with self.subTest(tool=tool):
                self.target.write_bytes(original)
                self.source.write_bytes(original)
                book = load_workbook(self.source)
                book["売上"]["A1"].font = Font(italic=True, color="0000FF")
                book.save(self.source)
                book.close()
                opened = []

                def picker(directory, merge=False):
                    opened.append("lazygit")
                    self.assertTrue(merge)
                    state = json.loads((directory / "session.json").read_text(encoding="utf-8"))
                    self.assertEqual((state["differences"], state["format_differences"], state["selectable"]), (0, 1, 1))
                    self.select(directory, state, formats={("売上", "A1")})
                    self.assertEqual(self.stage.apply_merge_session(directory), (0, 1))
                    return 0

                def launch(command, *args, **kwargs):
                    if command[0] != "WinMergeU.exe":
                        return original_run(command, *args, **kwargs)
                    opened.append("winmerge")
                    directory = Path(command[-2]).parent.parent
                    state = json.loads((directory / "session.json").read_text(encoding="utf-8"))
                    self.assertEqual((state["differences"], state["format_differences"], state["selectable"]), (0, 1, 1))
                    left, _, _ = self.win_paths(directory, state)
                    lines = self.copied_lines(directory, state, (), {("売上", "A1")})
                    left.write_bytes(("\n".join(lines) + "\n").encode("utf-8-sig"))
                    return subprocess.CompletedProcess(command, 0)

                output, errors = io.StringIO(), io.StringIO()
                with mock.patch.object(self.stage, "choose_tool", return_value=tool), mock.patch.object(
                    self.stage, "check_tools"
                ), mock.patch.object(self.stage, "run_picker", side_effect=picker), mock.patch.object(
                    self.stage, "find_winmerge", return_value="WinMergeU.exe"
                ), mock.patch.object(self.stage.subprocess, "run", side_effect=launch), mock.patch.object(
                    self.stage.sys, "stdout", output
                ), mock.patch.object(self.stage.sys, "stderr", errors):
                    status = self.stage.main(["--merge", str(self.target), str(self.source)])
                self.assertEqual(status, 0, errors.getvalue())
                self.assertEqual(opened, [tool])
                self.assertIn("書式 1 セルを取り込み先へ書き込みました:", output.getvalue())
                book = load_workbook(self.target)
                self.assertEqual(book["売上"]["A1"].value, "元の値")
                self.assertTrue(book["売上"]["A1"].font.italic)
                self.assertEqual(book["売上"]["A1"].font.color.rgb, "000000FF")
                book.close()

    def test_value_and_format_of_the_same_cell_can_be_imported_independently(self):
        original = self.target.read_bytes()
        book = load_workbook(self.source)
        book["売上"]["A1"].font = Font(italic=True, color="0000FF")
        book.save(self.source)
        book.close()
        with zipfile.ZipFile(self.target) as target, zipfile.ZipFile(self.source) as source:
            self.assertNotEqual(target.read("xl/styles.xml"), source.read("xl/styles.xml"))
        for tool in ("lazygit", "winmerge"):
            for values, formats in ((True, False), (False, True), (True, True)):
                with self.subTest(tool=tool, values=values, formats=formats):
                    self.target.write_bytes(original)
                    directory, state = self.prepare(tool)
                    value_keys = {("売上", "A1")} if values else set()
                    format_keys = {("売上", "A1")} if formats else set()
                    if tool == "lazygit":
                        self.select(directory, state, value_keys, format_keys)
                    else:
                        left, _, _ = self.win_paths(directory, state)
                        lines = self.copied_lines(directory, state, value_keys, format_keys)
                        left.write_bytes(("\n".join(lines) + "\n").encode("utf-8-sig"))
                    self.assertEqual(self.stage.apply_merge_session(directory), (int(values), int(formats)))
                    applied = json.loads((directory / "session.json").read_text(encoding="utf-8"))
                    self.assertEqual((applied["count"], applied["format_count"]), (int(values), int(formats)))
                    book = load_workbook(self.target)
                    cell = book["売上"]["A1"]
                    self.assertEqual(cell.value, "新しい値" if values else "元の値")
                    self.assertEqual(cell.font.italic, formats)
                    self.assertEqual(cell.font.bold, not formats)
                    self.assertEqual(book["売上"]["B1"].value, 10)
                    self.assertEqual(cell.fill.fgColor.rgb, "00FFFF00")
                    book.close()

    def test_winmerge_format_sections_follow_values_and_have_aligned_rows(self):
        book = load_workbook(self.source)
        book["売上"]["A1"].font = Font(italic=True)
        book.save(self.source)
        book.close()
        directory, state = self.prepare("winmerge")
        left, right, _ = self.win_paths(directory, state)
        before = left.read_text(encoding="utf-8-sig").splitlines()
        after = right.read_text(encoding="utf-8-sig").splitlines()
        self.assertEqual(len(before), len(after))
        heading = '# 書式 シート 1: "売上"'
        index = before.index(heading)
        self.assertEqual(after[index], heading)
        self.assertTrue(all(json.loads(line)[1] != "format" for line in before[3:index]))
        self.assertEqual(json.loads(before[index + 1])[:2], ["A1", "format"])
        self.assertEqual(before[index + 2], '# シート 2: "備考"')
        self.assertFalse(any(line.startswith("# 書式 シート 2:") for line in before))
        self.assertEqual((state["differences"], state["format_differences"], state["selectable"]), (5, 1, 6))

    def test_winmerge_invalid_format_copies_are_rejected(self):
        book = load_workbook(self.source)
        book["売上"]["A1"].font = Font(italic=True)
        book.save(self.source)
        book.close()
        for change in ("value-section", "description", "removed-heading"):
            with self.subTest(change=change):
                directory, state = self.prepare("winmerge")
                left, right, _ = self.win_paths(directory, state)
                lines = left.read_text(encoding="utf-8-sig").splitlines()
                changed = right.read_text(encoding="utf-8-sig").splitlines()
                heading = next(index for index, line in enumerate(lines) if line.startswith("# 書式 シート "))
                if change == "value-section":
                    first = next(index for index, line in enumerate(lines) if line.startswith("["))
                    lines[first] = changed[heading + 1]
                elif change == "description":
                    record = json.loads(changed[heading + 1])
                    record[2] += "、手で打った説明"
                    lines[heading + 1] = json.dumps(record, ensure_ascii=False, separators=(",", ":"))
                else:
                    del lines[heading]
                data = ("\n".join(lines) + "\n").encode("utf-8-sig")
                self.assert_unchanged_failure(lambda: self.run_win(directory, state, lambda _: data))

    def test_winmerge_format_records_are_separated_by_unchanged_blank_lines(self):
        book = load_workbook(self.source)
        book["売上"]["A1"].font = Font(italic=True)
        book["売上"]["Z9"].fill = PatternFill("solid", fgColor="FFFFFF00")
        book.save(self.source)
        book.close()
        for change in ("copied-into-blank", "removed-blank"):
            with self.subTest(change=change):
                directory, state = self.prepare("winmerge")
                left, right, _ = self.win_paths(directory, state)
                lines = left.read_text(encoding="utf-8-sig").splitlines()
                changed = right.read_text(encoding="utf-8-sig").splitlines()
                heading = next(index for index, line in enumerate(lines) if line.startswith("# 書式 シート "))
                self.assertEqual([line[:6] for line in lines[heading + 1:heading + 4]], ['["A1",', "", '["Z9",'])
                self.assertEqual(changed[heading + 2], "")
                if change == "copied-into-blank":
                    lines[heading + 2] = changed[heading + 1]
                else:
                    del lines[heading + 2]
                data = ("\n".join(lines) + "\n").encode("utf-8-sig")
                self.assert_unchanged_failure(lambda: self.run_win(directory, state, lambda _: data))
        directory, state = self.prepare("winmerge")
        chosen, _ = self.run_win(directory, state, lambda _: ("\n".join(self.copied_lines(directory, state, set(), {("売上", "Z9")})) + "\n").encode("utf-8-sig"))
        self.assertEqual(chosen, (set(), {(state["sheets"][0]["id"], "Z9")}))
        self.assertEqual(self.stage.apply_merge_session(directory), (0, 1))
        result = load_workbook(self.target)
        # An existing fill that differs only in alpha is reused.
        self.assertEqual(result["売上"]["Z9"].fill.fgColor.rgb[-6:], "FFFF00")
        self.assertFalse(result["売上"]["A1"].font.italic)
        result.close()

    def test_unsupported_formats_are_hidden_and_reasons_are_combined_per_cell(self):
        book = load_workbook(self.source)
        book["売上"]["A1"].font = Font(italic=True)
        book.save(self.source)
        book.close()
        rich_cell(self.source)
        with zipfile.ZipFile(self.source) as archive:
            styles = ET.fromstring(archive.read("xl/styles.xml"))
            sheet = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        index = int(sheet.find(".//{" + NS + "}c[@r='A1']").get("s"))
        xf = styles.find("{" + NS + "}cellXfs")[index]
        ET.SubElement(xf, "{" + NS + "}extLst")
        self.source.write_bytes(replace_parts(self.source.read_bytes(), {
            "xl/styles.xml": ET.tostring(styles, encoding="utf-8")
        }))
        for tool in ("lazygit", "winmerge"):
            with self.subTest(tool=tool):
                directory, state = self.prepare(tool)
                reasons = [item for item in state["unsupported"] if item["coordinate"] == "A1"]
                self.assertEqual(len(reasons), 1)
                self.assertIn("リッチテキスト", reasons[0]["reason"])
                self.assertIn("・", reasons[0]["reason"])
                self.assertIn("書式", reasons[0]["reason"])
                self.assertEqual((state["differences"], state["format_differences"], state["selectable"]), (5, 1, 4))
                if tool == "lazygit":
                    sheet = state["sheets"][0]
                    before = self.git("show", "HEAD:" + sheet["formats_filename"], cwd=directory / "picker").decode()
                    after = (directory / "picker" / sheet["formats_filename"]).read_text(encoding="utf-8")
                else:
                    left, right, _ = self.win_paths(directory, state)
                    before = next(line for line in left.read_text(encoding="utf-8-sig").splitlines()
                                  if line.startswith('["A1","format",'))
                    after = next(line for line in right.read_text(encoding="utf-8-sig").splitlines()
                                 if line.startswith('["A1","format",'))
                self.assertEqual(before, after)

    def test_unsupported_value_does_not_prevent_importing_its_cell_format(self):
        original = self.target.read_bytes()
        self.source.write_bytes(original)
        book = load_workbook(self.source)
        book["売上"]["A1"].font = Font(italic=True)
        book.save(self.source)
        book.close()
        rich_cell(self.source)
        for tool in ("lazygit", "winmerge"):
            with self.subTest(tool=tool):
                self.target.write_bytes(original)
                directory, state = self.prepare(tool)
                self.assertEqual((state["differences"], state["format_differences"], state["selectable"]), (1, 1, 1))
                if tool == "lazygit":
                    self.select(directory, state, formats={("売上", "A1")})
                else:
                    left, _, _ = self.win_paths(directory, state)
                    lines = self.copied_lines(directory, state, (), {("売上", "A1")})
                    left.write_bytes(("\n".join(lines) + "\n").encode("utf-8-sig"))
                self.assertEqual(self.stage.apply_merge_session(directory), (0, 1))
                book = load_workbook(self.target)
                self.assertEqual(book["売上"]["A1"].value, "元の値")
                self.assertTrue(book["売上"]["A1"].font.italic)
                book.close()

    def test_blank_cells_and_inherited_formats_can_be_imported(self):
        book = load_workbook(self.target)
        book["売上"].row_dimensions[4].font = Font(bold=True)
        book.save(self.target)
        book.close()
        original = self.target.read_bytes()
        self.source.write_bytes(original)
        book = load_workbook(self.source)
        book["売上"]["A4"] = "追加"
        book["売上"]["H3"].fill = PatternFill("solid", fgColor="00FF00")
        book.save(self.source)
        book.close()
        for tool in ("lazygit", "winmerge"):
            with self.subTest(tool=tool):
                self.target.write_bytes(original)
                directory, state = self.prepare(tool)
                self.assertEqual((state["differences"], state["format_differences"]), (1, 2))
                formats = {("売上", "A4"), ("売上", "H3")}
                if tool == "lazygit":
                    self.select(directory, state, formats=formats)
                else:
                    left, _, _ = self.win_paths(directory, state)
                    lines = self.copied_lines(directory, state, (), formats)
                    left.write_bytes(("\n".join(lines) + "\n").encode("utf-8-sig"))
                self.assertEqual(self.stage.apply_merge_session(directory), (0, 2))
                book = load_workbook(self.target)
                self.assertIsNone(book["売上"]["A4"].value)
                self.assertFalse(book["売上"]["A4"].font.bold)
                self.assertTrue(book["売上"].row_dimensions[4].font.bold)
                self.assertIsNone(book["売上"]["H3"].value)
                self.assertEqual(book["売上"]["H3"].fill.fgColor.rgb, "0000FF00")
                book.close()

    def test_same_file_and_normalized_path_are_rejected(self):
        (self.directory / "sub").mkdir()
        for source in (self.target, self.directory / "sub" / ".." / self.target.name):
            with self.subTest(source=source):
                self.assert_unchanged_failure(lambda: self.stage.prepare_merge_session(
                    self.target, source, self.directory / "bad", tool="lazygit"))

    def test_braces_in_sheet_name_keep_invalid_selection_a_stage_error(self):
        name = "売上 {年度} {coordinate}"
        for path in (self.target, self.source):
            book = load_workbook(path)
            book["売上"].title = name
            book.save(path)
            book.close()
        for tool in ("lazygit", "winmerge"):
            with self.subTest(tool=tool):
                directory, state = self.prepare(tool)
                if tool == "lazygit":
                    sheet = next(sheet for sheet in state["sheets"] if sheet["name"] == name)
                    path = directory / "picker" / sheet["filename"]
                    lines = path.read_text(encoding="utf-8").splitlines()
                    first = 0
                else:
                    path, _, _ = self.win_paths(directory, state)
                    lines = path.read_text(encoding="utf-8-sig").splitlines()
                    first = next(index for index, line in enumerate(lines) if line.startswith("["))
                record = json.loads(lines[first])
                record[2] = "手で打った値"
                lines[first] = json.dumps(record, ensure_ascii=False)
                path.write_text("\n".join(lines) + "\n", encoding="utf-8")
                if tool == "lazygit":
                    self.git("add", "--", sheet["filename"], cwd=directory / "picker")
                before = self.target.read_bytes()
                with self.assertRaises(self.stage.StageError) as raised:
                    self.stage.apply_merge_session(directory)
                self.assertIn(name + "!A1", str(raised.exception))
                self.assertEqual(self.target.read_bytes(), before)

    def test_zero_inode_same_file_detection_uses_normalized_real_paths(self):
        (self.directory / "sub").mkdir()
        alternate = self.directory / "sub" / ".." / self.target.name
        with mock.patch.object(type(self.target), "stat", return_value=types.SimpleNamespace(st_ino=0)), mock.patch.object(
            self.stage.os.path, "samefile", side_effect=AssertionError("zero inodes must use real paths")
        ) as samefile:
            self.assertTrue(self.stage._same_file(self.target, alternate))
            self.assertFalse(self.stage._same_file(self.target, self.source))
            samefile.assert_not_called()

    def test_hardlink_is_rejected(self):
        alias = self.directory / "hardlink.xlsx"
        try:
            os.link(self.target, alias)
        except OSError as error:
            self.skipTest(str(error))
        self.assert_unchanged_failure(lambda: self.stage.prepare_merge_session(
            self.target, alias, self.directory / "bad", tool="lazygit"))
        self.assert_unchanged_failure(lambda: self.prepare())

    @unittest.skipUnless(os.name == "posix", "POSIX symlinks do not require Windows privileges")
    def test_symlink_is_rejected(self):
        alias = self.directory / "symlink.xlsx"
        alias.symlink_to(self.target)
        self.assert_unchanged_failure(lambda: self.stage.prepare_merge_session(
            alias, self.source, self.directory / "bad", tool="lazygit"))
        self.assert_unchanged_failure(lambda: self.stage.prepare_merge_session(
            self.target, alias, self.directory / "bad", tool="lazygit"))

    @unittest.skipUnless(os.name == "nt", "Windows paths are case insensitive")
    def test_same_file_with_different_path_case_is_rejected(self):
        self.assert_unchanged_failure(lambda: self.stage.prepare_merge_session(
            self.target, str(self.target).upper(), self.directory / "bad", tool="lazygit"))

    def test_unsupported_extensions_on_either_side_are_rejected(self):
        for side in ("target", "source"):
            for suffix in (".xls", ".txt"):
                with self.subTest(side=side, suffix=suffix):
                    other = self.directory / (side + suffix)
                    other.write_bytes(getattr(self, side).read_bytes())
                    target, source = (other, self.source) if side == "target" else (self.target, other)
                    self.assert_unchanged_failure(lambda: self.stage.prepare_merge_session(
                        target, source, self.directory / "bad", tool="lazygit"), r"xlsx.*xlsm")

    def test_missing_directory_and_readonly_target_are_rejected_before_picker(self):
        for target in (self.directory / "missing.xlsx", self.directory / "folder.xlsx"):
            if target.name.startswith("folder"):
                target.mkdir()
            status, _, errors, _, lazy, win = self.main([str(target), str(self.source)])
            self.assertEqual(status, 1)
            self.assertIn("Excel の取り込みに失敗しました:", errors)
            lazy.assert_not_called()
            win.assert_not_called()
        with mock.patch.object(self.stage.os, "access", return_value=False):
            self.assert_unchanged_failure(lambda: self.prepare())

    def test_different_structure_reports_the_differing_item(self):
        book = load_workbook(self.source)
        book["備考"].title = "違う名前"
        book.save(self.source)
        book.close()
        self.assert_unchanged_failure(lambda: self.prepare(), "シート")

    def test_xlsx_and_xlsm_pair_is_accepted(self):
        macro = self.source.with_suffix(".xlsm")
        self.source.rename(macro)
        self.source = macro
        _, state = self.prepare()
        self.assertEqual(state["differences"], 5)

    def test_real_xlam_imports_only_selected_cells_through_git_alias(self):
        fixtures = ROOT / "tests" / "fixtures"
        self.target = self.directory / "取り込み先.xlam"
        self.source = self.directory / "取り込み元.xlam"
        self.target.write_bytes((fixtures / "sales.xlam").read_bytes())
        self.source.write_bytes((fixtures / "sales-edited.xlam").read_bytes())
        baseline, source = self.target.read_bytes(), self.source.read_bytes()
        directory, state = self.prepare()
        self.select(directory, state, {("売上", "B2")})
        alias = "!" + " ".join(shlex.quote(value) for value in
                               (Path(sys.executable).as_posix(), SCRIPT.as_posix()))
        self.git("-c", "alias.excel-merge=" + alias,
                 "excel-merge", "--apply-merge-session", str(directory))
        book = self.stage.Workbook(self.target.read_bytes())
        sheet = book.sheets[0].id
        self.assertEqual(book.cells[(sheet, "B2")], self.stage.Cell("number", "12"))
        self.assertEqual(book.cells[(sheet, "C3")], self.stage.Cell("number", "80"))
        self.assertEqual(self.source.read_bytes(), source)
        with zipfile.ZipFile(io.BytesIO(baseline)) as original, zipfile.ZipFile(self.target) as result:
            original_types = {node.get("PartName"): node.get("ContentType")
                              for node in ET.fromstring(original.read("[Content_Types].xml"))}
            result_types = {node.get("PartName"): node.get("ContentType")
                            for node in ET.fromstring(result.read("[Content_Types].xml"))}
            self.assertEqual(result_types["/xl/workbook.xml"], original_types["/xl/workbook.xml"])
            self.assertEqual(result_types["/xl/workbook.xml"],
                             "application/vnd.ms-excel.addin.macroEnabled.main+xml")

    def test_xlsb_lazygit_and_winmerge_import_the_same_selected_values(self):
        fixtures = ROOT / "tests" / "fixtures"
        baseline = (fixtures / "sales.xlsb").read_bytes()
        changed = (fixtures / "sales-edited.xlsb").read_bytes()
        selection = {("売上", "B2"), ("売上", "D4"), ("売上", "E5"),
                     ("売上", "F3"), ("売上", "A6"), ("売上", "AMK3"),
                     ("売上", "A50"), ("備考", "B2")}
        self.target = self.directory / "取り込み先 と空白.xlsb"
        self.source = self.directory / "取り込み元 と空白.xlsb"
        results = []
        for tool in ("lazygit", "winmerge"):
            with self.subTest(tool=tool):
                self.target.write_bytes(baseline)
                self.source.write_bytes(changed)
                directory, state = self.prepare(tool)
                self.assertEqual(state["workbook"], "xlsb")
                self.assertEqual((directory / "base.xlsb").read_bytes(), baseline)
                self.assertEqual((directory / "work.xlsb").read_bytes(), changed)
                self.assertEqual(state["format_differences"], 3)
                if tool == "lazygit":
                    self.select(directory, state, selection)
                else:
                    lines = self.copied_lines(directory, state, selection)
                    chosen, _ = self.run_win(directory, state, lambda _: ("\n".join(lines) + "\n").encode("utf-8-sig"))
                    self.assertEqual(len(chosen.values), len(selection))
                    self.assertEqual(chosen.formats, set())
                self.assertEqual(self.stage.apply_merge_session(directory), (len(selection), 0))
                base = self.stage.workbook_from_data(baseline)
                source = self.stage.workbook_from_data(changed)
                result = self.stage.workbook_from_data(self.target.read_bytes())
                self.assertEqual(result.workbook_format, "xlsb")
                expected = dict(base.cells)
                for name, coordinate in selection:
                    key = (next(sheet.id for sheet in base.sheets if sheet.name == name), coordinate)
                    if key in source.cells:
                        expected[key] = source.cells[key]
                    else:
                        expected.pop(key, None)
                    self.assertEqual(result.effective_style(key), base.effective_style(key))
                self.assertEqual(result.cells, expected)
                self.assertEqual(self.source.read_bytes(), changed)
                results.append(self.target.read_bytes())
        self.assertEqual(results[0], results[1])

    def test_xlsb_values_and_formats_are_independent_in_both_merge_tools(self):
        fixtures = ROOT / "tests" / "fixtures"
        baseline = (fixtures / "sales.xlsb").read_bytes()
        changed = (fixtures / "sales-edited.xlsb").read_bytes()
        self.target = self.directory / "書式取り込み先.xlsb"
        self.source = self.directory / "書式取り込み元.xlsb"
        self.source.write_bytes(changed)
        base = self.stage.workbook_from_data(baseline)
        source = self.stage.workbook_from_data(changed)
        key = (base.sheets[0].id, "B2")
        results = {}
        for tool in ("lazygit", "winmerge"):
            for values, formats in ((True, False), (False, True), (True, True)):
                with self.subTest(tool=tool, values=values, formats=formats):
                    self.target.write_bytes(baseline)
                    directory, state = self.prepare(tool)
                    selected = {("売上", "B2")}
                    value_keys = selected if values else set()
                    format_keys = selected if formats else set()
                    if tool == "lazygit":
                        for sheet in state["sheets"]:
                            self.assertTrue(sheet["formats_filename"].endswith(".formats.txt"))
                        self.select(directory, state, value_keys, format_keys)
                    else:
                        left, right, _ = self.win_paths(directory, state)
                        before = left.read_text(encoding="utf-8-sig").splitlines()
                        after = right.read_text(encoding="utf-8-sig").splitlines()
                        self.assertEqual(len(before), len(after))
                        heading = '# 書式 シート 1: "売上"'
                        self.assertEqual(after[before.index(heading)], heading)
                        lines = self.copied_lines(directory, state, value_keys, format_keys)
                        chosen, _ = self.run_win(directory, state,
                            lambda _: ("\n".join(lines) + "\n").encode("utf-8-sig"))
                        self.assertEqual((len(chosen.values), len(chosen.formats)), (int(values), int(formats)))
                    self.assertEqual(self.stage.apply_merge_session(directory), (int(values), int(formats)))
                    applied = json.loads((directory / "session.json").read_text(encoding="utf-8"))
                    self.assertEqual((applied["count"], applied["format_count"]), (int(values), int(formats)))
                    result = self.stage.workbook_from_data(self.target.read_bytes())
                    self.assertEqual(result.cells[key], (source if values else base).cells[key])
                    self.assertEqual({k: v for k, v in result.cells.items() if k != key},
                                     {k: v for k, v in base.cells.items() if k != key})
                    for cell_key in set(base.cell_styles) | set(source.cell_styles) | set(base.cells):
                        expected = source if formats and cell_key == key else base
                        self.assertEqual(result.stylesheet().exact(result.effective_style(cell_key)),
                                         expected.stylesheet().exact(expected.effective_style(cell_key)), cell_key)
                    self.assertEqual(self.source.read_bytes(), changed)
                    results[(tool, values, formats)] = self.target.read_bytes()
        for values, formats in ((True, False), (False, True), (True, True)):
            self.assertEqual(results[("lazygit", values, formats)], results[("winmerge", values, formats)])

    def test_xlsb_fill_and_blank_cell_formats_import_from_both_merge_tools(self):
        fixtures = ROOT / "tests" / "fixtures"
        baseline = (fixtures / "sales.xlsb").read_bytes()
        changed = (fixtures / "sales-edited.xlsb").read_bytes()
        self.target = self.directory / "空セル取り込み先.xlsb"
        self.source = self.directory / "空セル取り込み元.xlsb"
        self.source.write_bytes(changed)
        base = self.stage.workbook_from_data(baseline)
        source = self.stage.workbook_from_data(changed)
        results = {}
        for tool in ("lazygit", "winmerge"):
            for coordinates in (("A1",), ("M3",), ("A1", "M3")):
                with self.subTest(tool=tool, coordinates=coordinates):
                    self.target.write_bytes(baseline)
                    directory, state = self.prepare(tool)
                    selected = {("売上", coordinate) for coordinate in coordinates}
                    if tool == "lazygit":
                        self.select(directory, state, formats=selected)
                    else:
                        lines = self.copied_lines(directory, state, (), selected)
                        chosen, _ = self.run_win(directory, state,
                            lambda _: ("\n".join(lines) + "\n").encode("utf-8-sig"))
                        self.assertEqual((len(chosen.values), len(chosen.formats)), (0, len(coordinates)))
                    self.assertEqual(self.stage.apply_merge_session(directory), (0, len(coordinates)))
                    result = self.stage.workbook_from_data(self.target.read_bytes())
                    self.assertEqual(result.cells, base.cells)
                    self.assertNotIn((base.sheets[0].id, "M3"), result.cells)
                    for key in set(base.cell_styles) | set(source.cell_styles) | set(base.cells):
                        expected = source if key[0] == base.sheets[0].id and key[1] in coordinates else base
                        self.assertEqual(result.stylesheet().exact(result.effective_style(key)),
                                         expected.stylesheet().exact(expected.effective_style(key)), key)
                    self.assertEqual(self.source.read_bytes(), changed)
                    results[(tool, coordinates)] = self.target.read_bytes()
        for coordinates in (("A1",), ("M3",), ("A1", "M3")):
            self.assertEqual(results[("lazygit", coordinates)], results[("winmerge", coordinates)])

    def test_xlsb_and_xlsx_import_is_rejected_in_either_direction(self):
        fixtures = ROOT / "tests" / "fixtures"
        message = "形式の違うブック（xlsb と xlsx など）は取り込めません。同じ形式のブックを 2 つ指定してください。"
        for target_format, source_format in (("xlsb", "xlsx"), ("xlsx", "xlsb")):
            for tool in ("lazygit", "winmerge"):
                with self.subTest(target=target_format, source=source_format, tool=tool):
                    self.target = self.directory / ("混合先." + target_format)
                    self.source = self.directory / ("混合元." + source_format)
                    self.target.write_bytes((fixtures / ("sales." + target_format)).read_bytes())
                    self.source.write_bytes((fixtures / ("sales-edited." + source_format)).read_bytes())
                    source_before = self.source.read_bytes()
                    self.assert_unchanged_failure(lambda: self.prepare(tool), message)
                    self.assertEqual(self.source.read_bytes(), source_before)

    def test_mixed_workbook_contents_are_rejected_even_with_matching_extensions(self):
        fixtures = ROOT / "tests" / "fixtures"
        self.target.write_bytes((fixtures / "sales.xlsb").read_bytes())
        self.source.write_bytes((fixtures / "sales-edited.xlsx").read_bytes())
        self.assert_unchanged_failure(lambda: self.prepare(), "形式の違うブック")

    def test_unsupported_cells_on_either_side_are_hidden_from_source_projection(self):
        for side in ("target", "source"):
            with self.subTest(side=side):
                original = getattr(self, side).read_bytes()
                rich_cell(getattr(self, side))
                directory, state = self.prepare()
                self.assertTrue(any(item["sheet"] == "売上" and item["coordinate"] == "A1" and item["reason"]
                                    for item in state["unsupported"]))
                sheet = next(sheet for sheet in state["sheets"] if sheet["name"] == "売上")
                base = self.git("show", "HEAD:" + sheet["filename"], cwd=directory / "picker").decode().splitlines()
                work = (directory / "picker" / sheet["filename"]).read_text(encoding="utf-8").splitlines()
                self.assertEqual(next(line for line in base if json.loads(line)[0] == "A1"),
                                 next(line for line in work if json.loads(line)[0] == "A1"))
                getattr(self, side).write_bytes(original)

    def test_only_unsupported_differences_do_not_open_a_picker(self):
        self.source.write_bytes(self.target.read_bytes())
        rich_cell(self.source)
        status, output, errors, _, lazy, win = self.main([str(self.target), str(self.source)])
        self.assertEqual(status, 0, errors)
        self.assertIn("取り込めない違い（1 セル）: 売上!A1", output)
        self.assertIn("取り込める差分がありません（取り込めない違い 1 セル）。取り込み先は変えていません。", output)
        lazy.assert_not_called()
        win.assert_not_called()

    def test_selected_cells_preserve_unselected_values_styles_macros_and_source(self):
        marker = b"opaque macro bytes"
        self.target.write_bytes(replace_parts(self.target.read_bytes(), {"xl/vbaProject.bin": marker}))
        source = self.source.read_bytes(), self.source.stat().st_mtime_ns
        directory, state = self.prepare()
        self.select(directory, state, {("売上", "A1"), ("売上", "E1"), ("売上", "G1"), ("備考", "B4")})
        self.assertEqual(self.stage.apply_merge_session(directory), (4, 0))
        book = load_workbook(self.target)
        self.assertEqual(book["売上"]["A1"].value, "新しい値")
        self.assertEqual(book["売上"]["B1"].value, 10)
        self.assertIsNone(book["売上"]["E1"].value)
        self.assertEqual(book["売上"]["G1"].value, "追加")
        self.assertTrue(book["売上"]["A1"].font.bold)
        self.assertEqual(book["売上"]["A1"].fill.fgColor.rgb, "00FFFF00")
        book.close()
        with zipfile.ZipFile(self.target) as archive:
            self.assertEqual(archive.read("xl/vbaProject.bin"), marker)
            self.assertEqual(ET.fromstring(archive.read("xl/workbook.xml")).find("{" + NS + "}calcPr").get("fullCalcOnLoad"), "1")
        self.assertEqual((self.source.read_bytes(), self.source.stat().st_mtime_ns), source)
        self.assertTrue(json.loads((directory / "session.json").read_text(encoding="utf-8"))["applied"])

    def test_no_selection_preserves_target_bytes_and_timestamp(self):
        before = self.target.read_bytes(), self.target.stat().st_mtime_ns
        directory, _ = self.prepare()
        self.assertEqual(self.stage.apply_merge_session(directory), (0, 0))
        self.assertEqual((self.target.read_bytes(), self.target.stat().st_mtime_ns), before)
        self.assertFalse(json.loads((directory / "session.json").read_text(encoding="utf-8"))["applied"])

    def test_second_apply_and_wrong_session_targets_are_rejected(self):
        directory, state = self.prepare()
        self.select(directory, state, {("売上", "A1")})
        self.assert_unchanged_failure(lambda: self.stage.apply_session(directory))
        self.assertEqual(self.stage.apply_merge_session(directory), (1, 0))
        self.assert_unchanged_failure(lambda: self.stage.apply_merge_session(directory))
        directory, state = self.prepare()
        state["target"] = "index"
        self.stage.write_json(directory / "session.json", state)
        self.assert_unchanged_failure(lambda: self.stage.apply_merge_session(directory))

    def test_malformed_lazygit_selection_is_rejected(self):
        directory, state = self.prepare()
        sheet = state["sheets"][0]
        path = directory / "picker" / sheet["filename"]
        lines = path.read_text(encoding="utf-8").splitlines()
        record = json.loads(lines[0])
        record[2] = "手で作った値"
        lines[0] = json.dumps(record, ensure_ascii=False)
        path.write_text("\n".join(lines) + "\n", encoding="utf-8")
        self.git("add", "--", sheet["filename"], cwd=directory / "picker")
        self.assert_unchanged_failure(lambda: self.stage.apply_merge_session(directory), "A1")

    def test_changed_or_deleted_target_stops_without_leaving_a_temporary_file(self):
        for deleted in (False, True):
            with self.subTest(deleted=deleted):
                save_book(self.target)
                directory, state = self.prepare()
                self.select(directory, state, {("売上", "A1")})
                if deleted:
                    self.target.unlink()
                else:
                    self.edit(self.target, {("売上", "B1"): 999})
                files = set(self.directory.iterdir())
                with self.assertRaises((ValueError, OSError)):
                    self.stage.apply_merge_session(directory)
                self.assertEqual(set(self.directory.iterdir()), files)
                self.assertFalse(json.loads((directory / "session.json").read_text(encoding="utf-8"))["applied"])

    def test_source_is_the_opening_snapshot(self):
        directory, state = self.prepare()
        self.select(directory, state, {("売上", "A1")})
        self.edit(self.source, {("売上", "A1"): "後から変更"})
        self.assertEqual(self.stage.apply_merge_session(directory), (1, 0))
        book = load_workbook(self.target)
        self.assertEqual(book["売上"]["A1"].value, "新しい値")
        book.close()

    def test_unicode_record_separators_are_escaped_and_can_be_imported(self):
        value = "前\u2028中\x85後\u2029終"
        self.edit(self.source, {("売上", "A1"): value})
        for tool in ("lazygit", "winmerge"):
            with self.subTest(tool=tool):
                save_book(self.target)
                directory, state = self.prepare(tool)
                if tool == "lazygit":
                    texts = [(directory / "picker" / sheet["filename"]).read_bytes() for sheet in state["sheets"]]
                    self.select(directory, state, {("売上", "A1")})
                else:
                    left, right, _ = self.win_paths(directory, state)
                    texts = [left.read_bytes(), right.read_bytes()]
                    left.write_bytes(("\n".join(self.copied_lines(directory, state, {("売上", "A1")})) + "\n").encode("utf-8-sig"))
                for text in texts:
                    for separator in ("\u2028", "\x85", "\u2029"):
                        self.assertNotIn(separator.encode(), text)
                self.assertEqual(self.stage.apply_merge_session(directory), (1, 0))
                book = load_workbook(self.target)
                self.assertEqual(book["売上"]["A1"].value, value)
                book.close()

    def test_selected_number_boolean_and_formula_values_are_imported(self):
        self.edit(self.source, {("売上", "B1"): 123.5, ("売上", "C1"): False, ("売上", "D1"): "=B1*3"})
        directory, state = self.prepare()
        self.select(directory, state, {("売上", "B1"), ("売上", "C1"), ("売上", "D1")})
        self.assertEqual(self.stage.apply_merge_session(directory), (3, 0))
        book = load_workbook(self.target)
        self.assertEqual(book["売上"]["B1"].value, 123.5)
        self.assertIs(book["売上"]["C1"].value, False)
        self.assertEqual(book["売上"]["D1"].value, "=B1*3")
        book.close()

    def test_merge_snapshot_and_old_session_tampering_are_rejected(self):
        for change in ("base.xlsx", "work.xlsx", "old"):
            with self.subTest(change=change):
                directory, state = self.prepare()
                self.select(directory, state, {("売上", "A1")})
                if change == "old":
                    del state["base_sha256"]
                    self.stage.write_json(directory / "session.json", state)
                    pattern = "選択画面の形式が更新されています。画面を閉じて選び直してください。"
                else:
                    self.edit(directory / change, {("売上", "B1"): 999})
                    pattern = "変更"
                self.assert_unchanged_failure(lambda: self.stage.apply_merge_session(directory), pattern)

    @unittest.skipUnless(os.name == "posix", "POSIX readonly permissions")
    def test_actual_readonly_target_is_rejected(self):
        self.target.chmod(0o444)
        try:
            self.assert_unchanged_failure(lambda: self.prepare(), "読み取り専用|書き込")
        finally:
            self.target.chmod(0o600)

    def test_many_unsupported_differences_are_limited_to_twenty_examples(self):
        book = load_workbook(self.target)
        for row in range(1, 24):
            book["売上"].cell(row, 1, "元の値")
        book.save(self.target)
        book.close()
        self.source.write_bytes(self.target.read_bytes())
        for row in range(1, 24):
            rich_cell(self.source, "A" + str(row))
        status, output, errors, _, lazy, _ = self.main([str(self.target), str(self.source)])
        self.assertEqual(status, 0, errors)
        self.assertIn("取り込めない違い（23 セル）:", output)
        self.assertIn("売上!A20", output)
        self.assertNotIn("売上!A21", output)
        self.assertIn("ほか 3 セル", output)
        lazy.assert_not_called()

    def test_winmerge_permission_error_can_be_retried_from_terminal(self):
        original_run, original_replace = subprocess.run, self.stage.replace_file
        replacements = []

        def launch(command, *args, **kwargs):
            if command[0] != "WinMergeU.exe":
                return original_run(command, *args, **kwargs)
            directory = Path(command[-2]).parent.parent
            state = json.loads((directory / "session.json").read_text(encoding="utf-8"))
            left, _, _ = self.win_paths(directory, state)
            left.write_bytes(("\n".join(self.copied_lines(directory, state, {("売上", "A1")})) + "\n").encode("utf-8-sig"))
            return subprocess.CompletedProcess(command, 0)

        def replace(*args):
            replacements.append(args)
            if len(replacements) == 1:
                raise PermissionError("busy")
            return original_replace(*args)

        output, errors = io.StringIO(), io.StringIO()
        with mock.patch.object(self.stage, "find_winmerge", return_value="WinMergeU.exe"), mock.patch.object(
            self.stage.subprocess, "run", side_effect=launch
        ), mock.patch.object(self.stage, "replace_file", side_effect=replace), mock.patch.object(
            self.stage.sys.stdin, "isatty", return_value=True
        ), mock.patch("builtins.input", return_value="") as prompt, mock.patch.object(
            self.stage.sys, "stdout", output
        ), mock.patch.object(self.stage.sys, "stderr", errors):
            status = self.stage.main(["--merge", "--tool", "winmerge", "--", str(self.target), str(self.source)])
        self.assertEqual(status, 0, errors.getvalue())
        self.assertEqual(len(replacements), 2)
        prompt.assert_called_once_with("取り込み先を閉じてから Enter で再試行します（q で中止）: ")
        self.assertIn("1 セルを取り込み先へ書き込みました:", output.getvalue())

    def test_merge_picker_configuration(self):
        directory, _ = self.prepare()
        config = json.loads(self.stage.picker_config(directory, merge=True).read_text(encoding="utf-8"))
        command = config["customCommands"][0]
        self.assertIn("--apply-merge-session", command["command"])
        self.assertEqual(command["description"], "選択セルを取り込み先へ書き込む (反映後 q で戻る)")
        self.assertEqual(command["outputTitle"], "Excel の取り込み")

    def test_winmerge_text_has_bom_lf_aligned_headers_and_named_files(self):
        directory, state = self.prepare("winmerge")
        left, right, baseline = self.win_paths(directory, state)
        self.assertEqual(left.relative_to(directory).as_posix(), "winmerge/取り込み先.cells.txt")
        self.assertEqual(right.relative_to(directory).as_posix(), "winmerge/取り込み元.cells.txt")
        self.assertEqual(baseline.relative_to(directory).as_posix(), "baseline/取り込み先.cells.txt")
        self.assertEqual(left.read_bytes(), baseline.read_bytes())
        texts = [path.read_bytes() for path in (left, right)]
        for text in texts:
            self.assertTrue(text.startswith(b"\xef\xbb\xbf"))
            self.assertNotIn(b"\r", text)
            self.assertTrue(text.endswith(b"\n"))
        rows = [text.decode("utf-8-sig").splitlines() for text in texts]
        self.assertEqual(len(rows[0]), len(rows[1]))
        headers = [line for line in rows[0] if line.startswith("#")]
        self.assertEqual(headers, [line for line in rows[1] if line.startswith("#")])
        self.assertEqual(headers[:2], ["# 右（取り込み元）から左（取り込み先）へ、取り込むセルの行をコピーする",
                                     "# 上書き保存（Ctrl+S）してから WinMerge を閉じると、取り込み先のブックに書き込む"])
        self.assertIn('# シート 1: "売上"', headers)
        self.assertIn('# シート 2: "備考"', headers)

    def test_winmerge_some_and_all_cell_copies(self):
        self.apply_win({("売上", "A1"), ("備考", "B4")})
        book = load_workbook(self.target)
        self.assertEqual(book["売上"]["A1"].value, "新しい値")
        self.assertEqual(book["売上"]["B1"].value, 10)
        book.close()
        save_book(self.target)
        self.apply_win()
        self.assertEqual(self.stage.Workbook(self.target.read_bytes()).cells, self.stage.Workbook(self.source.read_bytes()).cells)

    def test_winmerge_close_without_save_selects_nothing_and_warns_on_quick_exit(self):
        before = self.target.read_bytes(), self.target.stat().st_mtime_ns
        directory, state = self.prepare("winmerge")
        errors = io.StringIO()
        with mock.patch.object(self.stage.sys, "stderr", errors), mock.patch.object(self.stage.time, "monotonic", side_effect=[10, 11]):
            selected, _ = self.run_win(directory, state)
        self.assertEqual(selected, self.stage.Selection(set(), set()))
        self.assertIn("WinMerge がすぐ終了しました", errors.getvalue())
        self.assertEqual(self.stage.apply_merge_session(directory), (0, 0))
        self.assertEqual((self.target.read_bytes(), self.target.stat().st_mtime_ns), before)

    def test_winmerge_line_endings_and_optional_bom_are_accepted(self):
        for encoding, separator, final in (("utf-8-sig", "\r\n", "\r\n"),
                                           ("utf-8", "\n", "\n"),
                                           ("utf-8", "\n", ""),
                                           ("utf-8-sig", "\n", "\n\n")):
            with self.subTest(encoding=encoding, separator=repr(separator), final=repr(final)):
                save_book(self.target)
                self.apply_win({("売上", "A1")}, encoding, separator, final)
        save_book(self.target)
        directory, state = self.prepare("winmerge")
        lines = self.copied_lines(directory, state, {("売上", "A1")})
        data = "".join(line + ("\n" if index % 2 else "\r\n") for index, line in enumerate(lines)).encode("utf-8-sig")
        selected, _ = self.run_win(directory, state, lambda _: data)
        self.assertEqual((len(selected.values), len(selected.formats)), (1, 0))

    def test_winmerge_non_utf8_saves_are_rejected(self):
        for encoding in ("utf-16", "cp932"):
            with self.subTest(encoding=encoding):
                directory, state = self.prepare("winmerge")
                data = ("\n".join(self.copied_lines(directory, state)) + "\n").encode(encoding)
                self.assert_unchanged_failure(lambda: self.run_win(directory, state, lambda _: data), "UTF-8")

    def test_winmerge_modified_headers_and_invalid_cell_replacements_are_rejected(self):
        for change in ("header", "removed", "added", "reordered", "manual", "other-sheet"):
            with self.subTest(change=change):
                directory, state = self.prepare("winmerge")
                lines = self.copied_lines(directory, state)
                records = [index for index, line in enumerate(lines) if line.startswith("[")]
                first, second, last = records[0], records[1], records[-1]
                if change == "header":
                    lines[0] += "変更"
                elif change == "removed":
                    del lines[first]
                elif change == "added":
                    lines.insert(first, lines[first])
                elif change == "reordered":
                    lines[first], lines[second] = lines[second], lines[first]
                elif change == "other-sheet":
                    lines[first] = lines[last]
                else:
                    record = json.loads(lines[first])
                    record[2] = "手で打った値"
                    lines[first] = json.dumps(record, ensure_ascii=False)
                data = ("\n".join(lines) + "\n").encode("utf-8-sig")
                self.assert_unchanged_failure(lambda: self.run_win(directory, state, lambda _: data),
                                              None if change == "header" else "A1")

    def test_winmerge_baseline_tampering_is_rejected(self):
        directory, state = self.prepare("winmerge")
        _, _, baseline = self.win_paths(directory, state)
        baseline.write_bytes(baseline.read_bytes() + b"\n")
        self.assert_unchanged_failure(lambda: self.run_win(directory, state), "射影|控え|変更")

    def test_winmerge_arguments_keep_left_writable_right_readonly_and_avoid_persistent_settings(self):
        directory, state = self.prepare("winmerge")
        _, launched = self.run_win(directory, state)
        left, right, _ = self.win_paths(directory, state)
        self.assertEqual(launched.call_args.args[0], ["WinMergeU.exe", "/s-", "/e", "/u", "/fl", "/wr", "/t", "text",
                         "/dl", "取り込み先（書き換える）: " + self.target.name,
                         "/dr", "取り込み元: " + self.source.name, str(left), str(right)])
        self.assertFalse(any(argument.lower().startswith(("/cfg", "/ignore", "/cp"))
                             for argument in launched.call_args.args[0]))

    def test_winmerge_validation_retry_preserves_saved_left(self):
        directory, state = self.prepare("winmerge")
        left, _, _ = self.win_paths(directory, state)
        valid = ("\n".join(self.copied_lines(directory, state, {("売上", "A1")})) + "\n").encode("utf-8-sig")
        saved = []

        def launch(command, *args, **kwargs):
            saved.append(left.read_bytes())
            left.write_bytes(b"invalid" if len(saved) == 1 else valid)
            return subprocess.CompletedProcess(command, 0)

        with mock.patch.object(self.stage, "find_winmerge", return_value="WinMergeU.exe"), mock.patch.object(
            self.stage.subprocess, "run", side_effect=launch
        ), mock.patch.object(self.stage.sys.stdin, "isatty", return_value=True), mock.patch("builtins.input", return_value="y") as prompt:
            self.assertEqual(tuple(map(len, self.stage.run_winmerge(directory))), (1, 0))
        self.assertEqual(saved[1], b"invalid")
        prompt.assert_called_once_with("WinMerge を開き直して直しますか [y/N]: ")

    def test_main_reports_paths_applies_winmerge_without_checking_lazygit_and_cleans_up(self):
        sessions = []
        original_run = subprocess.run

        def run(command, *args, **kwargs):
            if command[0] != "WinMergeU.exe":
                return original_run(command, *args, **kwargs)
            directory = Path(command[-2]).parent.parent
            sessions.append(directory)
            state = json.loads((directory / "session.json").read_text(encoding="utf-8"))
            left, _, _ = self.win_paths(directory, state)
            left.write_bytes(("\n".join(self.copied_lines(directory, state, {("売上", "A1")})) + "\n").encode("utf-8-sig"))
            return subprocess.CompletedProcess(command, 0)

        output, errors = io.StringIO(), io.StringIO()
        with mock.patch.object(self.stage, "find_winmerge", return_value="WinMergeU.exe"), mock.patch.object(
            self.stage.subprocess, "run", side_effect=run
        ), mock.patch.object(self.stage, "check_tools") as tools, mock.patch.object(
            self.stage, "run_picker"
        ) as lazy, mock.patch.object(self.stage.sys, "stdout", output), mock.patch.object(
            self.stage.sys, "stderr", errors
        ), mock.patch.object(self.stage.sys, "stdin", io.StringIO()):
            status = self.stage.main(["--merge", "--tool", "winmerge", "--", str(self.target), str(self.source)])
        self.assertEqual(status, 0, errors.getvalue())
        self.assertIn("書き換えるファイル: " + str(self.target.resolve()), output.getvalue())
        self.assertIn("取り込み元: " + str(self.source.resolve()), output.getvalue())
        self.assertIn("1 セルを取り込み先へ書き込みました: " + str(self.target.resolve()), output.getvalue())
        self.assertIn("上書き保存（Ctrl+S）", output.getvalue())
        tools.assert_not_called()
        lazy.assert_not_called()
        self.assertTrue(sessions)
        self.assertTrue(all(not directory.exists() for directory in sessions))

    def test_main_lazygit_cancel_and_apply_messages(self):
        before = self.target.read_bytes()
        status, output, errors, tools, _, _ = self.main([str(self.target), str(self.source)], picker=lambda directory, merge=False: 0)
        self.assertEqual(status, 0, errors)
        self.assertIn("取り込むセルが選ばれていません。取り込み先は変えていません。", output)
        self.assertIn("Enter → Space", output)
        tools.assert_called_once()
        self.assertEqual(self.target.read_bytes(), before)

        def picker(directory, merge=False):
            self.assertTrue(merge)
            state = json.loads((directory / "session.json").read_text(encoding="utf-8"))
            self.select(directory, state, {("売上", "A1")})
            self.assertEqual(self.stage.apply_merge_session(directory), (1, 0))
            return 0

        status, output, errors, _, _, _ = self.main([str(self.target), str(self.source)], picker=picker)
        self.assertEqual(status, 0, errors)
        self.assertIn("1 セルを取り込み先へ書き込みました:", output)

    def test_main_failure_reports_reason_and_cleans_session(self):
        sessions = []

        def picker(directory, merge=False):
            sessions.append(directory)
            raise self.stage.StageError("選択の試験用エラー")

        before = self.target.read_bytes()
        status, _, errors, _, _, _ = self.main([str(self.target), str(self.source)], picker=picker)
        self.assertEqual(status, 1)
        self.assertIn("Excel の取り込みに失敗しました: 選択の試験用エラー", errors)
        self.assertEqual(self.target.read_bytes(), before)
        self.assertTrue(sessions)
        self.assertTrue(all(not directory.exists() for directory in sessions))

    def test_date_system_and_defined_name_differences_report_the_item(self):
        original = self.source.read_bytes()
        for change, expected in (("date", "日付基準"), ("name", "名前定義")):
            with self.subTest(change=change):
                self.source.write_bytes(original)
                with zipfile.ZipFile(self.source) as archive:
                    root = ET.fromstring(archive.read("xl/workbook.xml"))
                if change == "date":
                    root.find("{" + NS + "}workbookPr").set("date1904", "1")
                else:
                    names = root.find("{" + NS + "}definedNames")
                    if names is None:
                        names = ET.SubElement(root, "{" + NS + "}definedNames")
                    ET.SubElement(names, "{" + NS + "}definedName", {"name": "ImportFixture"}).text = "売上!$A$1"
                self.source.write_bytes(replace_parts(self.source.read_bytes(), {
                    "xl/workbook.xml": ET.tostring(root, encoding="utf-8")
                }))
                self.assert_unchanged_failure(lambda: self.prepare(), expected)

    def test_hidden_apply_merge_session_messages(self):
        directory, state = self.prepare()
        output = io.StringIO()
        with mock.patch.object(self.stage.sys, "stdout", output):
            self.assertEqual(self.stage.main(["--apply-merge-session", str(directory)]), 0)
        self.assertIn("セル変更が選択されていません。Enter・Space でハンクを選び、X で反映してください。", output.getvalue())
        self.select(directory, state, {("売上", "A1")})
        output = io.StringIO()
        with mock.patch.object(self.stage.sys, "stdout", output):
            self.assertEqual(self.stage.main(["--apply-merge-session", str(directory)]), 0)
        self.assertIn("1 セルを取り込み先へ書き込みました。q で元の画面へ戻ってください。", output.getvalue())

    def test_resolve_user_path_honors_git_prefix_and_absolute_paths(self):
        with mock.patch.object(self.stage.Path, "cwd", return_value=self.directory), mock.patch.dict(os.environ, {"GIT_PREFIX": "sub/"}):
            self.assertEqual(self.stage.resolve_user_path("book.xlsx"), self.directory / "sub" / "book.xlsx")
            self.assertEqual(self.stage.resolve_user_path(self.target), self.target)

    def test_choose_tool_and_missing_winmerge_guidance(self):
        for platform, found, expected in (("nt", "winmerge.exe", "winmerge"), ("nt", None, "lazygit"),
                                         ("posix", "winmerge.exe", "lazygit")):
            with self.subTest(platform=platform, found=found), mock.patch.object(self.stage.os, "name", platform), mock.patch.object(
                self.stage, "find_winmerge", return_value=found
            ):
                self.assertEqual(self.stage.choose_tool("auto"), expected)
                self.assertEqual(self.stage.choose_tool("lazygit"), "lazygit")
        with mock.patch.object(self.stage, "find_winmerge", return_value=None):
            with self.assertRaisesRegex(self.stage.StageError, "WinMerge.*導入"):
                self.stage.choose_tool("winmerge")

    def test_find_winmerge_registry_local_programfiles_then_path(self):
        local = self.directory / "local"
        program = self.directory / "program"
        candidates = [self.directory / "registry-user.exe", self.directory / "registry-machine.exe",
                      local / "Programs" / "WinMerge" / "WinMergeU.exe", program / "WinMerge" / "WinMergeU.exe",
                      self.directory / "path.exe"]
        for candidate in candidates:
            candidate.parent.mkdir(parents=True, exist_ok=True)
            candidate.write_bytes(b"fake executable")
        registry = {"user": str(candidates[0]), "machine": str(candidates[1])}
        fake_registry = types.SimpleNamespace(HKEY_CURRENT_USER="user", HKEY_LOCAL_MACHINE="machine", KEY_READ=1,
                                             KEY_WOW64_32KEY=0, KEY_WOW64_64KEY=0)
        fake_registry.OpenKey = lambda root, *args: contextlib.nullcontext(root)
        fake_registry.QueryValueEx = lambda root, name: (registry[root], 1)
        with mock.patch.dict(sys.modules, {"winreg": fake_registry}), mock.patch.object(self.stage.os, "name", "nt"), mock.patch.object(
            self.stage, "Path", type(self.target)
        ), mock.patch.dict(os.environ, {"LOCALAPPDATA": str(local), "ProgramFiles": str(program)}), mock.patch.object(
            self.stage.shutil, "which", side_effect=lambda name: str(candidates[-1]) if candidates[-1].exists() else None
        ):
            for expected in candidates:
                self.assertEqual(str(self.stage.find_winmerge()), str(expected))
                expected.unlink()
            self.assertIsNone(self.stage.find_winmerge())

    def test_replace_file_uses_target_directory_and_leaves_no_temporary_file(self):
        before = set(self.directory.iterdir())
        original_mkstemp = self.stage.tempfile.mkstemp
        with mock.patch.object(self.stage.tempfile, "mkstemp", wraps=original_mkstemp) as create:
            self.stage.replace_file(self.target, b"replacement bytes", hashlib.sha256(self.target.read_bytes()).hexdigest())
        self.assertEqual(Path(create.call_args.kwargs["dir"]), self.target.parent)
        self.assertEqual(self.target.read_bytes(), b"replacement bytes")
        self.assertEqual(set(self.directory.iterdir()), before)

    @unittest.skipUnless(os.name == "posix", "POSIX mode bits")
    def test_replace_file_preserves_posix_mode(self):
        self.target.chmod(0o640)
        self.stage.replace_file(self.target, b"replacement", hashlib.sha256(self.target.read_bytes()).hexdigest())
        self.assertEqual(stat.S_IMODE(self.target.stat().st_mode), 0o640)

    def test_replace_file_retries_permission_errors_and_cleans_up_on_failure(self):
        digest = hashlib.sha256(self.target.read_bytes()).hexdigest()
        before = set(self.directory.iterdir())
        with mock.patch.object(self.stage.os, "replace", side_effect=[PermissionError("busy"), PermissionError("busy"), None]) as replace, mock.patch.object(
            self.stage.time, "sleep"
        ) as sleep:
            self.stage.replace_file(self.target, b"replacement", digest)
        self.assertEqual(replace.call_count, 3)
        self.assertEqual(sleep.call_args_list, [mock.call(0.2), mock.call(0.2)])
        self.assertEqual(set(self.directory.iterdir()), before)
        with mock.patch.object(self.stage.os, "replace", side_effect=PermissionError("busy")) as replace, mock.patch.object(
            self.stage.time, "sleep"
        ):
            with self.assertRaisesRegex(PermissionError, "Excel|閉じ"):
                self.stage.replace_file(self.target, b"replacement", digest)
        self.assertGreaterEqual(replace.call_count, 5)
        self.assertLessEqual(replace.call_count, 6)
        self.assertEqual(set(self.directory.iterdir()), before)

    def test_replace_file_rechecks_digest_after_a_permission_error(self):
        digest = hashlib.sha256(self.target.read_bytes()).hexdigest()
        before = set(self.directory.iterdir())

        def busy(*args):
            self.target.write_bytes(b"changed by another process")
            raise PermissionError("busy")

        with mock.patch.object(self.stage.os, "replace", side_effect=busy) as replace, mock.patch.object(self.stage.time, "sleep"):
            with self.assertRaises(ValueError):
                self.stage.replace_file(self.target, b"replacement", digest)
        replace.assert_called_once()
        self.assertEqual(self.target.read_bytes(), b"changed by another process")
        self.assertEqual(set(self.directory.iterdir()), before)

    @unittest.skipUnless(os.name == "posix", "LibreOffice lock detection is required on Linux")
    def test_libreoffice_lock_file_is_rejected(self):
        (self.target.parent / (".~lock." + self.target.name + "#")).write_bytes(b"locked")
        self.assert_unchanged_failure(lambda: self.prepare(), "LibreOffice|ロック")

    @unittest.skipUnless(os.name == "nt", "Windows alternate data streams")
    def test_replace_file_preserves_zone_identifier(self):
        zone = Path(str(self.target) + ":Zone.Identifier")
        marker = b"[ZoneTransfer]\r\nZoneId=3\r\n"
        zone.write_bytes(marker)
        self.stage.replace_file(self.target, b"replacement", hashlib.sha256(self.target.read_bytes()).hexdigest())
        self.assertEqual(zone.read_bytes(), marker)

    @unittest.skipUnless(os.name == "nt", "Windows sharing violations")
    def test_open_target_is_rejected_before_reading_workbooks(self):
        import ctypes
        from ctypes import wintypes
        kernel = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                                      wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        kernel.CreateFileW.restype = wintypes.HANDLE
        kernel.CloseHandle.argtypes = [wintypes.HANDLE]
        handle = kernel.CreateFileW(str(self.target), 0x80000000, 1, None, 3, 0x80, None)
        self.assertNotEqual(handle, wintypes.HANDLE(-1).value)
        try:
            with mock.patch.object(self.stage, "Workbook") as workbook:
                self.assert_unchanged_failure(lambda: self.prepare(), "Excel|開")
            workbook.assert_not_called()
        finally:
            kernel.CloseHandle(handle)

    @unittest.skipUnless(os.name == "posix", "Git shell aliases are tested on POSIX")
    def test_alias_resolves_paths_from_subdirectory_and_outside_repository(self):
        spec = importlib.util.spec_from_file_location("excel_merge_setup_alias", ROOT / "scripts" / "setup-excel-diff.py")
        setup = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(setup)
        helper = self.directory / "merge wrapper.py"
        helper.write_text("import importlib.util,sys\n"
                          + "sys.path.insert(0," + repr(str(SCRIPT.parent)) + ")\n"
                          + "s=importlib.util.spec_from_file_location('merge_alias_helper'," + repr(str(SCRIPT)) + ")\n"
                          + "m=importlib.util.module_from_spec(s);s.loader.exec_module(m)\n"
                          + "m.check_tools=lambda:None\nm.choose_tool=lambda name:'lazygit'\n"
                          + "m.run_picker=lambda directory,merge=False:0\n"
                          + "raise SystemExit(m.main(sys.argv[1:]))\n", encoding="utf-8")
        env = dict(self.env, GIT_CONFIG_GLOBAL=str(self.directory / "gitconfig"))
        self.git("config", "--global", "alias.excel-merge", setup.merge_alias(sys.executable, helper), env=env)
        repo = self.directory / "repo"
        sub = repo / "sub"
        sub.mkdir(parents=True)
        self.git("init", "--quiet", cwd=repo, env=env)
        target, source = sub / "target.xlsx", sub / "source.xlsx"
        shutil.copyfile(self.target, target)
        shutil.copyfile(self.source, source)
        output = self.git("excel-merge", "--", target.name, source.name, cwd=sub, env=env).decode()
        self.assertIn("書き換えるファイル: " + str(target.resolve()), output)
        output = self.git("excel-merge", "--", self.target.name, self.source.name, env=env).decode()
        self.assertIn("書き換えるファイル: " + str(self.target.resolve()), output)


if __name__ == "__main__":
    unittest.main()
