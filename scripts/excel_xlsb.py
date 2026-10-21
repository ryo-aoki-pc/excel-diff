"""Read BIFF12 workbook records without changing the original package.

Record and cell offsets deliberately refer to uncompressed part bytes.  They
remain useful when a later edit replaces individual records in those parts.
"""

import datetime as dt
import hashlib
import io
import json
import math
import posixpath
import struct
import zipfile
from array import array
from contextlib import nullcontext
from dataclasses import dataclass, field
from urllib.parse import unquote

from lxml import etree

from excel_formula import FormulaError, decode_formula, quote_sheet_name


MAIN_TYPE = "application/vnd.ms-excel.sheet.binary.macroEnabled.main"
_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_CELL_TYPES = set(range(1, 12)) | {62}
_ERRORS = {0: "#NULL!", 7: "#DIV/0!", 15: "#VALUE!", 23: "#REF!", 29: "#NAME?", 36: "#NUM!", 42: "#N/A", 43: "#GETTING_DATA"}


class WorkbookError(ValueError):
    """A binary workbook cannot be read safely."""


@dataclass(frozen=True)
class Record:
    type: int
    start: int
    payload_start: int
    end: int
    payload: bytes


@dataclass(frozen=True)
class Cell:
    row: int
    col: int
    kind: str
    value: object
    ixfe: int
    start: int
    end: int
    record: Record
    rgce: bytes = b""
    rgcb: bytes = b""
    rich: bool = False

    @property
    def coordinate(self):
        return _coordinate(self.row, self.col)


@dataclass(frozen=True)
class Row:
    row: int
    ixfe: int
    spans: tuple
    record: Record


@dataclass(frozen=True)
class DefinedName:
    name: str
    scope: int
    rgce: bytes
    rgcb: bytes
    record: Record


@dataclass(frozen=True)
class FormulaRange:
    bounds: tuple
    rgce: bytes
    rgcb: bytes
    record: Record


@dataclass
class Sheet:
    id: str
    name: str
    path: str
    state: int
    rel_id: str
    records: list = field(default_factory=list)
    cells: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    arrays: dict = field(default_factory=dict)
    shared: dict = field(default_factory=dict)
    tables: dict = field(default_factory=dict)
    merges: list = field(default_factory=list)
    columns: list = field(default_factory=list)
    dimensions: object = None


def record_header(record_type, size):
    """Encode the variable-length BIFF12 record header (MS-XLSB 2.1.4)."""
    if not isinstance(record_type, int) or not 0 <= record_type <= 0x3FFF:
        raise WorkbookError("xlsb のレコードの種類が範囲外です")
    if not isinstance(size, int) or not 0 <= size <= 0x0FFFFFFF:
        raise WorkbookError("xlsb のレコードの長さが範囲外です")
    header = bytearray()
    if record_type < 128:
        header.append(record_type)
    else:
        header.extend(((record_type & 0x7F) | 0x80, record_type >> 7))
    while size >= 128:
        header.append((size & 0x7F) | 0x80)
        size >>= 7
    header.append(size)
    return bytes(header)


def records(data):
    """Yield bounded records, retaining header, payload, and end offsets."""
    position = 0
    while position < len(data):
        start = position
        first = data[position]
        position += 1
        record_type = first & 0x7F
        if first & 0x80:
            if position == len(data) or data[position] & 0x80:
                raise WorkbookError("xlsb のレコードの種類が壊れています")
            record_type |= data[position] << 7
            position += 1
        size = 0
        for shift in range(4):
            if position == len(data):
                raise WorkbookError("xlsb のレコードの長さが途中で切れています")
            value = data[position]
            position += 1
            size |= (value & 0x7F) << (shift * 7)
            if not value & 0x80:
                break
        else:
            raise WorkbookError("xlsb のレコードの長さが壊れています")
        end = position + size
        if end > len(data):
            raise WorkbookError("xlsb のレコードが途中で切れています")
        yield Record(record_type, start, position, end, data[position:end])
        position = end


class _RecordView:
    """Expose one record without retaining a copy of its payload."""

    __slots__ = ("_records", "_index")

    def __init__(self, records, index):
        self._records, self._index = records, index

    @property
    def type(self):
        return self._records._types[self._index]

    @property
    def start(self):
        return self._records._ends[self._index - 1] if self._index else 0

    @property
    def payload_start(self):
        return self._records._payloads[self._index]

    @property
    def end(self):
        return self._records._ends[self._index]

    @property
    def payload(self):
        return self._records.data[self.payload_start:self.end]


class _RecordSequence:
    """Store record positions in packed arrays; create short-lived views on access."""

    __slots__ = ("data", "_types", "_payloads", "_ends")

    def __init__(self, data):
        self.data = data
        offset_type = "I" if len(data) <= 0xFFFFFFFF else "Q"
        self._types, self._payloads, self._ends = array("H"), array(offset_type), array(offset_type)
        types, payloads, ends = self._types.append, self._payloads.append, self._ends.append
        length, position = len(data), 0
        while position < length:
            first = data[position]
            position += 1
            kind = first & 0x7F
            if first & 0x80:
                if position == length or data[position] & 0x80:
                    raise WorkbookError("xlsb のレコードの種類が壊れています")
                kind |= data[position] << 7
                position += 1
            size = 0
            for shift in range(0, 28, 7):
                if position == length:
                    raise WorkbookError("xlsb のレコードの長さが途中で切れています")
                value = data[position]
                position += 1
                size |= (value & 0x7F) << shift
                if not value & 0x80:
                    break
            else:
                raise WorkbookError("xlsb のレコードの長さが壊れています")
            end = position + size
            if end > length:
                raise WorkbookError("xlsb のレコードが途中で切れています")
            types(kind)
            payloads(position)
            ends(end)
            position = end

    def __len__(self):
        return len(self._types)

    def __getitem__(self, index):
        if isinstance(index, slice):
            return [self[item] for item in range(*index.indices(len(self)))]
        if index < 0:
            index += len(self)
        if not 0 <= index < len(self):
            raise IndexError(index)
        return _RecordView(self, index)

    def __iter__(self):
        for index in range(len(self)):
            yield _RecordView(self, index)


def _wide_at(data, position, end):
    if position + 4 > end:
        raise WorkbookError("xlsb のレコードの内容が途中で切れています")
    length = struct.unpack_from("<I", data, position)[0]
    finish = position + 4 + length * 2
    if finish > end:
        raise WorkbookError("xlsb のレコードの内容が途中で切れています")
    try:
        return data[position + 4:finish].decode("utf-16le"), finish
    except UnicodeError as error:
        raise WorkbookError("xlsb の文字列が壊れています") from error


class _CellView:
    """Decode a cell from the shared sheet bytes only when its value is needed."""

    __slots__ = ("_cells", "_index")

    def __init__(self, cells, index):
        self._cells, self._index = cells, index

    @property
    def row(self):
        return self._cells._rows[self._index]

    @property
    def record(self):
        return self._cells.records[self._cells._indices[self._index]]

    @property
    def col(self):
        records, index = self._cells.records, self._cells._indices[self._index]
        return struct.unpack_from("<I", records.data, records._payloads[index])[0]

    @property
    def ixfe(self):
        records, index = self._cells.records, self._cells._indices[self._index]
        return struct.unpack_from("<I", records.data, records._payloads[index] + 4)[0] & 0xFFFFFF

    @property
    def start(self):
        index = self._cells._indices[self._index]
        return self._cells.records._ends[index - 1] if index else 0

    @property
    def end(self):
        return self._cells.records._ends[self._cells._indices[self._index]]

    @property
    def coordinate(self):
        return _coordinate(self.row, self.col)

    @property
    def kind(self):
        kind = self._cells.records._types[self._cells._indices[self._index]]
        return _CELL_KINDS[kind]

    @property
    def value(self):
        return self.read()[1]

    @property
    def rich(self):
        records, index = self._cells.records, self._cells._indices[self._index]
        kind, position = records._types[index], records._payloads[index] + 8
        if kind == 7:
            return self._cells.strings[struct.unpack_from("<I", records.data, position)[0]][1]
        return kind == 62 and bool(records.data[position] & 1)

    @property
    def rgce(self):
        return self._formula()[0]

    @property
    def rgcb(self):
        return self._formula()[1]

    def _formula(self):
        records, index = self._cells.records, self._cells._indices[self._index]
        kind = records._types[index]
        if kind not in (8, 9, 10, 11):
            return b"", b""
        data, position, end = records.data, records._payloads[index] + 8, records._ends[index]
        if kind == 8:
            position += 4 + struct.unpack_from("<I", data, position)[0] * 2
        else:
            position += 8 if kind == 9 else 1
        return self._cells.formula_at(position + 2, end)

    def read(self):
        """Return kind, cached value, rich flag, rgce, and rgcb in one pass."""
        records, index = self._cells.records, self._cells._indices[self._index]
        return self._cells._read_at(records._types[index], records._payloads[index], records._ends[index])


_CELL_KINDS = {1: "blank", 2: "number", 3: "error", 4: "boolean", 5: "number",
               6: "string", 7: "string", 8: "string", 9: "number", 10: "boolean",
               11: "error", 62: "string"}


class _CellSequence:
    __slots__ = ("records", "strings", "_rows", "_indices")

    def __init__(self, records, strings):
        self.records, self.strings = records, strings
        index_type = "I" if len(records) <= 0xFFFFFFFF else "Q"
        self._rows, self._indices = array("I"), array(index_type)

    def __len__(self):
        return len(self._rows)

    def __getitem__(self, index):
        if isinstance(index, slice):
            return [self[item] for item in range(*index.indices(len(self)))]
        if index < 0:
            index += len(self)
        if not 0 <= index < len(self):
            raise IndexError(index)
        return _CellView(self, index)

    def __iter__(self):
        for index in range(len(self)):
            yield _CellView(self, index)

    def iter_values(self):
        """Yield decoded cell scalars without constructing cell or record views."""
        records, rows, indices = self.records, self._rows, self._indices
        data, types, payloads, ends = records.data, records._types, records._payloads, records._ends
        unpack, decode = struct.unpack_from, self._read_at
        for index, record_index in enumerate(indices):
            position, end, record_type = payloads[record_index], ends[record_index], types[record_index]
            col, flags = unpack("<II", data, position)
            kind, value, rich, rgce, rgcb = decode(record_type, position, end)
            start = ends[record_index - 1] if record_index else 0
            yield rows[index], col, flags & 0xFFFFFF, record_type, start, kind, value, rich, rgce, rgcb, index

    def _read_at(self, kind, position, end):
        data = self.records.data
        position += 8
        value, rich = None, False
        if kind == 2:
            number = struct.unpack_from("<I", data, position)[0]
            if number & 2:
                value = (number - 0x100000000 if number & 0x80000000 else number) >> 2
            else:
                value = struct.unpack("<d", struct.pack("<II", 0, number & 0xFFFFFFFC))[0]
            if number & 1:
                value /= 100
        elif kind in (5, 9):
            value = struct.unpack_from("<d", data, position)[0]
            position += 8
        elif kind in (3, 11):
            value = _ERRORS.get(data[position], "#VALUE!")
            position += 1
        elif kind in (4, 10):
            value = bool(data[position])
            position += 1
        elif kind in (6, 8):
            value, position = _wide_at(data, position, end)
        elif kind == 7:
            value, rich = self.strings[struct.unpack_from("<I", data, position)[0]]
        elif kind == 62:
            value, rich = _rich_string(_Reader(data[position:end]))
        rgce, rgcb = self.formula_at(position + 2, end) if kind in (8, 9, 10, 11) else (b"", b"")
        return _CELL_KINDS[kind], value, rich, rgce, rgcb

    def formula_at(self, position, end, copy=True):
        data = self.records.data
        if position + 4 > end:
            raise WorkbookError("xlsb のレコードの内容が途中で切れています")
        size = struct.unpack_from("<I", data, position)[0]
        start = position + 4
        position = start + size
        if position + 4 > end:
            raise WorkbookError("xlsb のレコードの内容が途中で切れています")
        extra = struct.unpack_from("<I", data, position)[0]
        if position + 4 + extra > end:
            raise WorkbookError("xlsb のレコードの内容が途中で切れています")
        if copy:
            return data[start:position], data[position + 4:position + 4 + extra]
        return None


class _RowView:
    __slots__ = ("_rows", "_index")

    def __init__(self, rows, index):
        self._rows, self._index = rows, index

    @property
    def record(self):
        return self._rows.records[self._rows._indices[self._index]]

    @property
    def row(self):
        record = self.record
        return struct.unpack_from("<I", record._records.data, record.payload_start)[0]

    @property
    def ixfe(self):
        record = self.record
        return struct.unpack_from("<I", record._records.data, record.payload_start + 4)[0]

    @property
    def spans(self):
        record = self.record
        data, position = record._records.data, record.payload_start
        count = struct.unpack_from("<I", data, position + 13)[0]
        return tuple(struct.iter_unpack("<II", memoryview(data)[position + 17:position + 17 + count * 8]))


class _RowSequence:
    __slots__ = ("records", "_indices")

    def __init__(self, records):
        index_type = "I" if len(records) <= 0xFFFFFFFF else "Q"
        self.records, self._indices = records, array(index_type)

    def __len__(self):
        return len(self._indices)

    def __getitem__(self, index):
        if isinstance(index, slice):
            return [self[item] for item in range(*index.indices(len(self)))]
        if index < 0:
            index += len(self)
        if not 0 <= index < len(self):
            raise IndexError(index)
        return _RowView(self, index)

    def __iter__(self):
        for index in range(len(self)):
            yield _RowView(self, index)


class _PartsArchive:
    def __init__(self, parts):
        self.parts = parts

    def read(self, path):
        return self.parts[path]


class _Reader:
    def __init__(self, data, position=0):
        self.data = data
        self.position = position

    def take(self, size):
        end = self.position + size
        if size < 0 or end > len(self.data):
            raise WorkbookError("xlsb のレコードの内容が途中で切れています")
        value = self.data[self.position:end]
        self.position = end
        return value

    def unpack(self, fmt):
        return struct.unpack("<" + fmt, self.take(struct.calcsize("<" + fmt)))

    def uint(self):
        return self.unpack("I")[0]

    def wide(self, nullable=False):
        length = self.uint()
        if nullable and length == 0xFFFFFFFF:
            return None
        try:
            return self.take(length * 2).decode("utf-16le")
        except UnicodeError as error:
            raise WorkbookError("xlsb の文字列が壊れています") from error

    def formula(self):
        rgce = self.take(self.uint())
        rgcb = self.take(self.uint())
        return rgce, rgcb


def _xml(data, description):
    try:
        parser = etree.XMLParser(resolve_entities=False, no_network=True)
        tree = etree.parse(io.BytesIO(data), parser)
        if tree.docinfo.doctype:
            raise WorkbookError("DTD を含む XML は扱えません: " + description)
        return tree.getroot()
    except etree.XMLSyntaxError as error:
        raise WorkbookError("xlsb の XML を読めません: " + description) from error


def _target(owner, target):
    target = unquote(target.split("#", 1)[0])
    path = posixpath.normpath(posixpath.join(posixpath.dirname(owner), target))
    if target.startswith("/"):
        path = posixpath.normpath(target.lstrip("/"))
    if path in ("", ".", "..") or path.startswith("../") or "\\" in path:
        raise WorkbookError("xlsb の部品の参照が不正です")
    return path


def _relationships(archive, owner):
    path = posixpath.join(posixpath.dirname(owner), "_rels", posixpath.basename(owner) + ".rels") if owner else "_rels/.rels"
    root = _xml(archive.read(path), path)
    result = {}
    for node in root:
        if node.tag != "{" + _REL_NS + "}Relationship":
            continue
        key = node.get("Id")
        if not key or key in result:
            raise WorkbookError("xlsb の関連付けが重複しています")
        external = node.get("TargetMode", "").casefold() == "external"
        result[key] = (node.get("Type", ""), None if external else _target(owner, node.get("Target", "")))
    return result


def workbook_path(archive):
    office = [target for kind, target in _relationships(archive, "").values() if kind.endswith("/officeDocument") and target]
    if len(office) != 1:
        raise WorkbookError("xlsb のブックの関連付けが見つかりません")
    path = office[0]
    types = _xml(archive.read("[Content_Types].xml"), "[Content_Types].xml")
    default_type = None
    for node in types:
        if node.tag.endswith("}Override") and node.get("PartName", "").lstrip("/") == path:
            return path if node.get("ContentType") == MAIN_TYPE else None
        if node.tag.endswith("}Default") and node.get("Extension", "").casefold() == path.rpartition(".")[2].casefold():
            default_type = node.get("ContentType")
    return path if default_type == MAIN_TYPE else None


def is_xlsb(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            return workbook_path(archive) is not None
    except (KeyError, ValueError, zipfile.BadZipFile):
        return False


def _rich_string(reader):
    flags = reader.take(1)[0]
    value = reader.wide()
    if flags & 1:
        reader.take(reader.uint() * 4)
    if flags & 2:
        reader.wide()
        reader.take(reader.uint() * 10)
    return value, bool(flags & 1)


def _bounds(reader):
    bounds = reader.unpack("IIII")
    first_row, last_row, first_col, last_col = bounds
    if not 0 <= first_row <= last_row < 1048576 or not 0 <= first_col <= last_col < 16384:
        raise WorkbookError("xlsb のセル範囲が範囲外です")
    return bounds


def _coordinate(row, col):
    letters = ""
    col += 1
    while col:
        col, remainder = divmod(col - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters + str(row + 1)


def _reference(bounds):
    first_row, last_row, first_col, last_col = bounds
    return _coordinate(first_row, first_col) + ":" + _coordinate(last_row, last_col)


def _rk(data):
    number = struct.unpack("<I", data)[0]
    if number & 2:
        signed = number - 0x100000000 if number & 0x80000000 else number
        value = signed >> 2
    else:
        value = struct.unpack("<d", struct.pack("<II", 0, number & 0xFFFFFFFC))[0]
    return value / 100 if number & 1 else value


class BinaryWorkbook:
    """Read-only package model; all positions are in the original part bytes."""

    def __init__(self, data, *, parts=None, sheet_paths=None, compact=False):
        self._initialize(data)
        self._compact = compact
        if parts is not None:
            self.parts = parts
        try:
            opened = zipfile.ZipFile(io.BytesIO(data)) if parts is None else nullcontext(_PartsArchive(parts))
            with opened as archive:
                self._read_metadata(archive)
                for sheet in self.sheets:
                    if sheet.path and (sheet_paths is None or sheet.path in sheet_paths):
                        self._read_sheet(sheet, self._part(archive, sheet.path))
        except (KeyError, zipfile.BadZipFile, EOFError, struct.error) as error:
            raise WorkbookError("xlsb のブックが壊れているため読めません: " + str(error)) from error

    def _initialize(self, data):
        self.data = data
        self._compact = False
        self.sheets = []
        self.names = []
        self.supporting_links = []
        self.xti = []
        self.strings = []
        self.number_formats = {}
        self.cell_formats = []
        self.date1904 = False
        self.calc_record = None
        self.parts = {}

    def _read_metadata(self, archive):
        self.path = workbook_path(archive)
        if self.path is None:
            raise WorkbookError("xlsb のブックではありません")
        self.relationships = _relationships(archive, self.path)
        self.workbook_records = list(records(self._part(archive, self.path)))
        self._read_workbook()
        for kind, path in self.relationships.values():
            if path and kind.endswith("/sharedStrings"):
                self._read_strings(self._part(archive, path))
            elif path and kind.endswith("/styles"):
                self._read_styles(self._part(archive, path))

    def _part(self, archive, path):
        if path not in self.parts:
            self.parts[path] = archive.read(path)
        return self.parts[path]

    def _read_workbook(self):
        link = None
        for record in self.workbook_records:
            reader = _Reader(record.payload)
            if record.type == 156:
                state, sheet_id = reader.unpack("II")
                rel_id = reader.wide(nullable=True)
                name = reader.wide()
                relation = self.relationships.get(rel_id)
                if relation is None and rel_id is not None:
                    raise WorkbookError("xlsb のシートの関連付けが見つかりません")
                path = relation[1] if relation else None
                self.sheets.append(Sheet(str(sheet_id), name, path, state, rel_id))
            elif record.type == 153:
                self.date1904 = bool(reader.uint() & 1)
            elif record.type == 39:
                reader.uint()
                reader.take(1)
                scope = reader.uint()
                name = reader.wide()
                rgce, rgcb = reader.formula()
                self.names.append(DefinedName(name, scope, rgce, rgcb, record))
            elif record.type in (355, 357, 358, 667):
                link = {"type": record.type, "names": [], "record": record}
                self.supporting_links.append(link)
            elif record.type == 361 and link is not None:
                link["names"].append(reader.wide())
            elif record.type == 362:
                for _ in range(reader.uint()):
                    self.xti.append(reader.unpack("Iii"))
            elif record.type == 157:
                self.calc_record = record

    def _read_strings(self, data):
        for record in records(data):
            if record.type == 19:
                self.strings.append(_rich_string(_Reader(record.payload)))

    def _read_styles(self, data):
        in_cell_xfs = False
        for record in records(data):
            reader = _Reader(record.payload)
            if record.type == 44:
                index = reader.unpack("H")[0]
                self.number_formats[index] = reader.wide()
            elif record.type == 617:
                in_cell_xfs = True
            elif record.type == 618:
                in_cell_xfs = False
            elif record.type == 47 and in_cell_xfs:
                reader.take(2)
                self.cell_formats.append(reader.unpack("H")[0])

    def _read_sheet(self, sheet, data):
        if self._compact:
            return self._read_sheet_compact(sheet, data)
        row = None
        last_formula = None
        seen = set()
        sheet.records = list(records(data))
        for record in sheet.records:
            reader = _Reader(record.payload)
            if record.type == 0:
                row, ixfe = reader.unpack("II")
                if row >= 1048576:
                    raise WorkbookError("xlsb の行番号が範囲外です")
                reader.take(5)
                spans = tuple(reader.unpack("II") for _ in range(reader.uint()))
                if any(not 0 <= first <= last < 16384 for first, last in spans):
                    raise WorkbookError("xlsb の行の列範囲が不正です")
                sheet.rows.append(Row(row, ixfe, spans, record))
            elif record.type in _CELL_TYPES:
                if row is None:
                    raise WorkbookError("xlsb のセルの前に行の記録がありません")
                col, flags = reader.unpack("II")
                if col >= 16384 or (row, col) in seen:
                    raise WorkbookError("xlsb のセル番地が範囲外か重複しています")
                seen.add((row, col))
                kind, value, rich = "blank", None, False
                if record.type == 2:
                    kind, value = "number", _rk(reader.take(4))
                elif record.type in (3, 11):
                    kind, value = "error", _ERRORS.get(reader.take(1)[0], "#VALUE!")
                elif record.type in (4, 10):
                    kind, value = "boolean", bool(reader.take(1)[0])
                elif record.type in (5, 9):
                    kind, value = "number", reader.unpack("d")[0]
                elif record.type in (6, 8):
                    kind, value = "string", reader.wide()
                elif record.type == 7:
                    kind = "string"
                    index = reader.uint()
                    if index >= len(self.strings):
                        raise WorkbookError("xlsb の共有文字列の番号が範囲外です")
                    value, rich = self.strings[index]
                elif record.type == 62:
                    kind = "string"
                    value, rich = _rich_string(reader)
                rgce, rgcb = b"", b""
                if record.type in (8, 9, 10, 11):
                    reader.take(2)
                    rgce, rgcb = reader.formula()
                    last_formula = (row, col)
                if kind == "number" and not math.isfinite(value):
                    raise WorkbookError("xlsb のセルの数値が有限ではありません")
                sheet.cells.append(Cell(row, col, kind, value, flags & 0xFFFFFF, record.start, record.end, record, rgce, rgcb, rich))
            elif record.type in (426, 427):
                bounds = _bounds(reader)
                if record.type == 426:
                    reader.take(1)
                rgce, rgcb = reader.formula()
                target = sheet.arrays if record.type == 426 else sheet.shared
                anchor = (bounds[0], bounds[2]) if record.type == 426 else last_formula
                if anchor is None or not (bounds[0] <= anchor[0] <= bounds[1] and bounds[2] <= anchor[1] <= bounds[3]):
                    raise WorkbookError("xlsb の共有数式の基準セルが不正です")
                target[anchor] = FormulaRange(bounds, rgce, rgcb, record)
            elif record.type == 428:
                bounds = _bounds(reader)
                inputs = reader.unpack("IIII")
                flags = reader.take(1)[0]
                table = {"ref": _reference(bounds), "dt2D": bool(flags & 2), "dtr": bool(flags & 1), "r1": None if flags & 4 else _coordinate(inputs[0], inputs[1]), "r2": _coordinate(inputs[2], inputs[3]) if flags & 2 and not flags & 8 else None, "del1": bool(flags & 4), "del2": bool(flags & 8), "ca": bool(flags & 16)}
                sheet.tables[(bounds[0], bounds[2])] = (bounds, table, record)
            elif record.type == 176:
                sheet.merges.append((_bounds(reader), record))
            elif record.type == 60:
                first, last, width, ixfe = reader.unpack("IIII")
                sheet.columns.append((first, last, width, ixfe, record))
            elif record.type == 148:
                sheet.dimensions = _bounds(reader)
        sheet.cells.sort(key=lambda cell: (cell.row, cell.col))

    def _read_sheet_compact(self, sheet, data):
        records = sheet.records = _RecordSequence(data)
        cells = sheet.cells = _CellSequence(records, self.strings)
        rows = sheet.rows = _RowSequence(records)
        payloads, ends, types = records._payloads, records._ends, records._types
        unpack = struct.unpack_from
        row = last_formula = None
        previous, unsorted = -1, False
        seen = set()
        for index, kind in enumerate(types):
            position, end = payloads[index], ends[index]
            if kind == 0:
                if end - position < 17:
                    raise WorkbookError("xlsb のレコードの内容が途中で切れています")
                row, ixfe = unpack("<II", data, position)
                if row >= 1048576:
                    raise WorkbookError("xlsb の行番号が範囲外です")
                if rows._indices and row <= previous:
                    raise WorkbookError("xlsb の行の記録が不正です")
                previous = row
                seen.clear()
                count = unpack("<I", data, position + 13)[0]
                if position + 17 + count * 8 > end:
                    raise WorkbookError("xlsb のレコードの内容が途中で切れています")
                for first, last in struct.iter_unpack("<II", memoryview(data)[position + 17:position + 17 + count * 8]):
                    if not 0 <= first <= last < 16384:
                        raise WorkbookError("xlsb の行の列範囲が不正です")
                rows._indices.append(index)
            elif kind in _CELL_TYPES:
                if row is None:
                    raise WorkbookError("xlsb のセルの前に行の記録がありません")
                if end - position < 8:
                    raise WorkbookError("xlsb のレコードの内容が途中で切れています")
                col = unpack("<I", data, position)[0]
                if col >= 16384 or col in seen:
                    raise WorkbookError("xlsb のセル番地が範囲外か重複しています")
                if seen and col < previous_col:
                    unsorted = True
                seen.add(col)
                previous_col = col
                value_position = position + 8
                if kind in (2, 7):
                    minimum = value_position + 4
                elif kind in (3, 4, 10, 11):
                    minimum = value_position + 1
                elif kind in (5, 9):
                    minimum = value_position + 8
                elif kind in (6, 8):
                    if value_position + 4 > end:
                        raise WorkbookError("xlsb のレコードの内容が途中で切れています")
                    minimum = value_position + 4 + unpack("<I", data, value_position)[0] * 2
                elif kind == 62:
                    _rich_string(_Reader(data[value_position:end]))
                    minimum = end
                else:
                    minimum = value_position
                if minimum > end:
                    raise WorkbookError("xlsb のレコードの内容が途中で切れています")
                if kind == 7 and unpack("<I", data, value_position)[0] >= len(self.strings):
                    raise WorkbookError("xlsb の共有文字列の番号が範囲外です")
                if kind in (5, 9) and not math.isfinite(unpack("<d", data, value_position)[0]):
                    raise WorkbookError("xlsb のセルの数値が有限ではありません")
                if kind == 2:
                    number = unpack("<I", data, value_position)[0]
                    if not number & 2 and not math.isfinite(struct.unpack("<d", struct.pack("<II", 0, number & 0xFFFFFFFC))[0]):
                        raise WorkbookError("xlsb のセルの数値が有限ではありません")
                if kind in (6, 8):
                    _wide_at(data, value_position, end)
                if kind in (8, 9, 10, 11):
                    cells.formula_at(minimum + 2, end, copy=False)
                    last_formula = (row, col)
                cells._rows.append(row)
                cells._indices.append(index)
            else:
                if kind not in (426, 427, 428, 176, 60, 148):
                    continue
                record = records[index]
                reader = _Reader(record.payload)
                if kind in (426, 427):
                    bounds = _bounds(reader)
                    if kind == 426:
                        reader.take(1)
                    rgce, rgcb = reader.formula()
                    target = sheet.arrays if kind == 426 else sheet.shared
                    anchor = (bounds[0], bounds[2]) if kind == 426 else last_formula
                    if anchor is None or not (bounds[0] <= anchor[0] <= bounds[1] and bounds[2] <= anchor[1] <= bounds[3]):
                        raise WorkbookError("xlsb の共有数式の基準セルが不正です")
                    target[anchor] = FormulaRange(bounds, rgce, rgcb, record)
                elif kind == 428:
                    bounds = _bounds(reader)
                    inputs = reader.unpack("IIII")
                    flags = reader.take(1)[0]
                    table = {"ref": _reference(bounds), "dt2D": bool(flags & 2), "dtr": bool(flags & 1), "r1": None if flags & 4 else _coordinate(inputs[0], inputs[1]), "r2": _coordinate(inputs[2], inputs[3]) if flags & 2 and not flags & 8 else None, "del1": bool(flags & 4), "del2": bool(flags & 8), "ca": bool(flags & 16)}
                    sheet.tables[(bounds[0], bounds[2])] = (bounds, table, record)
                elif kind == 176:
                    sheet.merges.append((_bounds(reader), record))
                elif kind == 60:
                    first, last, width, ixfe = reader.unpack("IIII")
                    sheet.columns.append((first, last, width, ixfe, record))
                elif kind == 148:
                    sheet.dimensions = _bounds(reader)
        if unsorted:
            order = sorted(range(len(cells)), key=lambda index: (cells._rows[index], unpack("<I", data, payloads[cells._indices[index]])[0]))
            cells._rows = array("I", (cells._rows[index] for index in order))
            cells._indices = array(cells._indices.typecode, (cells._indices[index] for index in order))

    def formula_context(self, sheet):
        sheets, external_names = {}, {}
        for index, (link_index, first, last) in enumerate(self.xti):
            if link_index >= len(self.supporting_links):
                continue
            link = self.supporting_links[link_index]
            if link["type"] == 357 and 0 <= first <= last < len(self.sheets):
                sheets[index] = self.sheets[first].name if first == last else (self.sheets[first].name, self.sheets[last].name)
            elif link["type"] == 358:
                sheets[index] = sheet.name
                for name_index, name in enumerate(link["names"], 1):
                    if any(item.name == name for item in self.names):
                        external_names[(index, name_index)] = name
            elif link["type"] == 667:
                for name_index, name in enumerate(link["names"], 1):
                    external_names[(index, name_index)] = name
        names = {}
        for index, defined in enumerate(self.names, 1):
            name = defined.name
            if defined.scope != 0xFFFFFFFF and defined.scope < len(self.sheets) and self.sheets[defined.scope] is not sheet:
                name = quote_sheet_name(self.sheets[defined.scope].name) + "!" + name
            names[index] = name
        return {"sheets": sheets, "names": names, "external_names": external_names}

    def cell_value(self, cell):
        if cell.kind != "number" or not self.cell_formats:
            return cell.value
        from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format
        from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel

        if cell.ixfe >= len(self.cell_formats):
            raise WorkbookError("xlsb のセル書式の番号が範囲外です")
        index = self.cell_formats[cell.ixfe]
        format_code = self.number_formats.get(index, BUILTIN_FORMATS.get(index, "General"))
        if not is_date_format(format_code):
            return cell.value
        try:
            return from_excel(cell.value, epoch=MAC_EPOCH if self.date1904 else WINDOWS_EPOCH, timedelta=is_timedelta_format(format_code))
        except (ValueError, OverflowError):
            return "#VALUE!"


def read_workbook(data):
    return BinaryWorkbook(data)


def sheet_titles(data):
    """Read sheet names without parsing cell parts, for large-file placeholders."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            path = workbook_path(archive)
            if path is None:
                raise WorkbookError("xlsb のブックではありません")
            result = []
            for record in records(archive.read(path)):
                if record.type == 156:
                    reader = _Reader(record.payload, 8)
                    reader.wide(nullable=True)
                    result.append(reader.wide())
            return result
    except (KeyError, zipfile.BadZipFile) as error:
        raise WorkbookError("xlsb のシート名を読めません") from error


def _quoted(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def _value_text(value):
    if isinstance(value, dt.datetime):
        return "datetime: " + value.isoformat()
    if isinstance(value, dt.date):
        return "date: " + value.isoformat()
    if isinstance(value, dt.time):
        return "time: " + value.isoformat()
    if isinstance(value, dt.timedelta):
        return "duration-seconds: " + _quoted(value.total_seconds())
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return _quoted(value)


def _stream_records(stream, include=None):
    """Read one bounded record at a time from a buffered ZIP member."""
    read = stream.read
    while True:
        first = read(1)
        if not first:
            return
        record_type = first[0] & 0x7F
        if first[0] & 0x80:
            second = read(1)
            if not second or second[0] & 0x80:
                raise WorkbookError("xlsb のレコードの種類が壊れています")
            record_type |= second[0] << 7
        size = 0
        for shift in range(4):
            byte = read(1)
            if not byte:
                raise WorkbookError("xlsb のレコードの長さが途中で切れています")
            size |= (byte[0] & 0x7F) << (shift * 7)
            if not byte[0] & 0x80:
                break
        else:
            raise WorkbookError("xlsb のレコードの長さが壊れています")
        # Discard uninterested payloads in bounded chunks during the metadata pass.
        keep = include is None or record_type in include
        if keep:
            payload = read(size)
            if len(payload) != size:
                raise WorkbookError("xlsb のレコードが途中で切れています")
            yield record_type, payload
        else:
            remaining = size
            while remaining:
                payload = read(min(remaining, 65536))
                if not payload:
                    raise WorkbookError("xlsb のレコードが途中で切れています")
                remaining -= len(payload)


def _data_table_text(table):
    from openpyxl.worksheet.formula import DataTableFormula

    properties = dict(table)
    # Excel writes both flags even when false; a two-input table is row-oriented.
    properties["dt2D"] = "1" if table["dt2D"] else "0"
    properties["dtr"] = "1" if table["dt2D"] or table["dtr"] else "0"
    return "data-table-formula: " + _quoted(dict(DataTableFormula(**properties)))


class _TextWorkbook(BinaryWorkbook):
    """Keep only workbook metadata and a single row of worksheet cell payloads."""

    def __init__(self, data, archive):
        self._initialize(data)
        self._read_metadata(archive)
        self._date_formats = {}

    def _part(self, archive, path):
        return archive.read(path)

    def _read_sheet_metadata(self, sheet, stream):
        # Definitions may follow a PtgExp reference in a later row.  Resolve them
        # before rendering, without retaining the surrounding cell records.
        row = last_formula = None
        include = {0, 8, 9, 10, 11, 426, 427, 428, 176, 60, 148}
        for record_type, payload in _stream_records(stream, include):
            reader = _Reader(payload)
            if record_type == 0:
                row = reader.uint()
                if row >= 1048576:
                    raise WorkbookError("xlsb の行番号が範囲外です")
                reader.take(9)
                spans = tuple(reader.unpack("II") for _ in range(reader.uint()))
                if any(not 0 <= first <= last < 16384 for first, last in spans):
                    raise WorkbookError("xlsb の行の列範囲が不正です")
            elif record_type in (8, 9, 10, 11):
                if row is None:
                    raise WorkbookError("xlsb のセルの前に行の記録がありません")
                last_formula = (row, reader.uint())
            elif record_type in (426, 427):
                bounds = _bounds(reader)
                if record_type == 426:
                    reader.take(1)
                rgce, rgcb = reader.formula()
                anchor = (bounds[0], bounds[2]) if record_type == 426 else last_formula
                if anchor is None or not (bounds[0] <= anchor[0] <= bounds[1] and bounds[2] <= anchor[1] <= bounds[3]):
                    raise WorkbookError("xlsb の共有数式の基準セルが不正です")
                target = sheet.arrays if record_type == 426 else sheet.shared
                target[anchor] = FormulaRange(bounds, rgce, rgcb, None)
            elif record_type == 428:
                bounds = _bounds(reader)
                inputs = reader.unpack("IIII")
                flags = reader.take(1)[0]
                table = {"ref": _reference(bounds), "dt2D": bool(flags & 2), "dtr": bool(flags & 1), "r1": None if flags & 4 else _coordinate(inputs[0], inputs[1]), "r2": _coordinate(inputs[2], inputs[3]) if flags & 2 and not flags & 8 else None, "del1": bool(flags & 4), "del2": bool(flags & 8), "ca": bool(flags & 16)}
                sheet.tables[(bounds[0], bounds[2])] = (bounds, table, None)
            elif record_type in (148, 176):
                _bounds(reader)
            elif record_type == 60:
                reader.unpack("IIII")

    def _cached_text(self, kind, value, ixfe):
        if kind == "error":
            return "error: " + _quoted(value)
        if kind == "number" and self.cell_formats:
            if ixfe >= len(self.cell_formats):
                raise WorkbookError("xlsb のセル書式の番号が範囲外です")
            if ixfe not in self._date_formats:
                from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format, is_timedelta_format

                index = self.cell_formats[ixfe]
                code = self.number_formats.get(index, BUILTIN_FORMATS.get(index, "General"))
                self._date_formats[ixfe] = (is_date_format(code), is_timedelta_format(code))
            is_date, is_delta = self._date_formats[ixfe]
            if is_date:
                from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel

                try:
                    value = from_excel(value, epoch=MAC_EPOCH if self.date1904 else WINDOWS_EPOCH, timedelta=is_delta)
                except (ValueError, OverflowError):
                    return 'error: "#VALUE!"'
        return _value_text(value)

    def _cell_text(self, row, record_type, payload, sheet, context):
        reader = _Reader(payload)
        col, flags = reader.unpack("II")
        kind, value = "blank", None
        if record_type == 2:
            kind, value = "number", _rk(reader.take(4))
        elif record_type in (3, 11):
            kind, value = "error", _ERRORS.get(reader.take(1)[0], "#VALUE!")
        elif record_type in (4, 10):
            kind, value = "boolean", bool(reader.take(1)[0])
        elif record_type in (5, 9):
            kind, value = "number", reader.unpack("d")[0]
        elif record_type in (6, 8):
            kind, value = "string", reader.wide()
        elif record_type == 7:
            kind = "string"
            index = reader.uint()
            if index >= len(self.strings):
                raise WorkbookError("xlsb の共有文字列の番号が範囲外です")
            value, _ = self.strings[index]
        elif record_type == 62:
            kind = "string"
            value, _ = _rich_string(reader)
        if kind == "number" and not math.isfinite(value):
            raise WorkbookError("xlsb のセルの数値が有限ではありません")
        if kind == "blank":
            return None
        text = self._cached_text(kind, value, flags & 0xFFFFFF)
        rgce, rgcb = b"", b""
        if record_type in (8, 9, 10, 11):
            reader.take(2)
            rgce, rgcb = reader.formula()
        if (row, col) in sheet.tables:
            return _data_table_text(sheet.tables[(row, col)][1])
        if record_type not in (8, 9, 10, 11):
            return text
        array = None
        try:
            if not rgce:
                raise FormulaError("数式のトークン列が空です")
            if rgce[0] == 1:
                if len(rgce) != 5 or len(rgcb) < 4:
                    raise FormulaError("配列・共有数式の参照が壊れています")
                anchor = (struct.unpack("<I", rgce[1:])[0], struct.unpack("<I", rgcb[:4])[0])
                array = sheet.arrays.get(anchor)
                shared = sheet.shared.get(anchor)
                if array:
                    if (row, col) != anchor:
                        return text
                    rgce, rgcb = array.rgce, array.rgcb
                elif shared:
                    rgce, rgcb = shared.rgce, shared.rgcb
                elif anchor in sheet.tables:
                    return _data_table_text(sheet.tables[anchor][1]) if (row, col) == anchor else text
                else:
                    raise FormulaError("配列・共有数式の式が見つかりません")
            formula = "=" + decode_formula(rgce, rgcb, version=12, row=row, col=col, **context)
            return "array-formula: " + _quoted({"text": formula, "ref": _reference(array.bounds)}) if array else "formula: " + _quoted(formula)
        except FormulaError:
            fingerprint = hashlib.sha256(rgce + rgcb).hexdigest()[:8]
            return 'formula-unparsed: "#' + fingerprint + '" ' + text

    def write_sheet(self, sheet, stream, out):
        row = None
        row_cells = []
        seen = set()
        context = self.formula_context(sheet)

        def write_row():
            for col, record_type, payload in sorted(row_cells):
                text = self._cell_text(row, record_type, payload, sheet, context)
                if text is not None:
                    out.write(_coordinate(row, col) + "\t" + text + "\n")

        for record_type, payload in _stream_records(stream, _CELL_TYPES | {0}):
            if record_type == 0:
                write_row()
                next_row = _Reader(payload).uint()
                if row is not None and next_row <= row:
                    raise WorkbookError("xlsb の行の記録の順序が不正です")
                row = next_row
                row_cells.clear()
                seen.clear()
            elif record_type in _CELL_TYPES:
                if row is None:
                    raise WorkbookError("xlsb のセルの前に行の記録がありません")
                col = _Reader(payload).uint()
                if col >= 16384 or col in seen:
                    raise WorkbookError("xlsb のセル番地が範囲外か重複しています")
                seen.add(col)
                row_cells.append((col, record_type, payload))
        write_row()


def render_text(data):
    out = io.StringIO()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            book = _TextWorkbook(data, archive)
            for sheet in book.sheets:
                out.write("Sheet: " + _quoted(sheet.name) + "\n")
                if sheet.path:
                    with archive.open(sheet.path) as member, io.BufferedReader(member) as stream:
                        book._read_sheet_metadata(sheet, stream)
                    with archive.open(sheet.path) as member, io.BufferedReader(member) as stream:
                        book.write_sheet(sheet, stream, out)
    except (KeyError, zipfile.BadZipFile, EOFError, struct.error) as error:
        raise WorkbookError("xlsb のブックが壊れているため読めません: " + str(error)) from error
    return out.getvalue()
