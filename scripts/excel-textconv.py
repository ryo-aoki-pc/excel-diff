#!/usr/bin/env python3
"""Render Excel cells and Word documents as stable UTF-8 text for Git's textconv diff driver."""

import argparse
import datetime as dt
import hashlib
import io
import json
import os
import sys
import time
from pathlib import Path

# Modules needed only to convert or to start the worker are imported where they
# are used: a cache hit runs twice for every Excel diff lazygit shows.

WORD_SUFFIXES = {".docx", ".docm", ".dotx", ".dotm"}
WORD_MAIN_TYPES = {
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml",
    "application/vnd.ms-word.document.macroEnabled.main+xml",
    "application/vnd.ms-word.template.macroEnabledTemplate.main+xml",
}

OLE_SIGNATURE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
# With --cache-dir, a conversion that takes longer than this many seconds is
# finished by the background worker so Git (and lazygit) never wait for it.
DEFAULT_TIMEOUT = 2.0
TIMEOUT_VARIABLE = "EXCEL_TEXTCONV_TIMEOUT"
# Twice the fastest rate measured (about 4 MB/s of sheet XML on a Ryzen AI
# MAX+ 395), so only workbooks that cannot finish in time skip the wait.
OPTIMISTIC_BYTES_PER_SECOND = 8 * 1024 * 1024
# The same rule for BIFF12: about twice the rate measured for unpacked sheet and
# string records (about 4.3 MB/s for 500,000 cells on the same machine).
XLSB_OPTIMISTIC_BYTES_PER_SECOND = 8 * 1024 * 1024
CACHE_LIMIT_BYTES = 512 * 1024 * 1024
CACHE_MAX_AGE = 30 * 24 * 3600
PRUNE_INTERVAL = 3600
# The worker refreshes its lock while it runs; an older lock belongs to a
# worker that died, so another one may start.
WORKER_HEARTBEAT = 5
WORKER_STALE = 60
# Requests nobody repeated for this long are dropped instead of converted.
REQUEST_MAX_AGE = 3600


def quoted(value):
    # Escape tabs/newlines so a cell always occupies exactly one diff line.
    return json.dumps(value, ensure_ascii=False, allow_nan=False)


def value_text(value):
    if isinstance(value, dt.datetime):
        return "datetime: " + value.isoformat()
    if isinstance(value, dt.date):
        return "date: " + value.isoformat()
    if isinstance(value, dt.time):
        return "time: " + value.isoformat()
    if isinstance(value, dt.timedelta):
        return "duration-seconds: " + quoted(value.total_seconds())
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return quoted(value)


def column_letter(index):
    # Local helper: importing openpyxl just for this costs ~0.25 s per .xls.
    letters = ""
    while index:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def render_ooxml(source, out):
    from openpyxl import load_workbook
    from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

    # Git's temporary blobs have no Excel extension. Passing an open stream also
    # lets openpyxl read .xlsm/.xltx/.xltm without relying on the filename.
    book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in book:
            out.write("Sheet: " + quoted(sheet.title) + "\n")
            # Read actual rows, including files with incorrect dimension metadata.
            sheet.reset_dimensions()
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    value = cell.value
                    if isinstance(value, ArrayFormula):
                        text = "array-formula: " + quoted(
                            {"text": value.text, "ref": value.ref}
                        )
                    elif isinstance(value, DataTableFormula):
                        text = "data-table-formula: " + quoted(dict(value))
                    elif cell.data_type == "f":
                        text = "formula: " + quoted(value)
                    elif cell.data_type == "e":
                        text = "error: " + quoted(value)
                    else:
                        text = value_text(value)
                    out.write(cell.coordinate + "\t" + text + "\n")
    finally:
        book.close()


def render_xls(data, out):
    try:
        out.write(sibling_module("excel_xls").render_text(data))
    except Exception as error:
        raise ValueError("xls/xlsb を読めません。パスワード付き（暗号化）の xlsb は表示できません: "
                         + str(error)) from error


def package_main(data):
    """Resolve the officeDocument part and its declared content type."""
    import posixpath
    from urllib.parse import unquote
    import xml.etree.ElementTree as ET
    import zipfile

    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ET.fromstring(archive.read("_rels/.rels"))
        office = [
            node for node in root
            if node.get("Type", "").endswith("/officeDocument")
            and node.get("TargetMode", "").lower() != "external"
        ]
        if len(office) != 1:
            return None, None
        target = unquote(office[0].get("Target", "").split("#", 1)[0])
        target = posixpath.normpath("/" + target.lstrip("/"))
        types = ET.fromstring(archive.read("[Content_Types].xml"))
        for node in types:
            if (node.tag.endswith("}Override")
                    and node.get("PartName", "").casefold() == target.casefold()):
                return target.lstrip("/"), node.get("ContentType")
        extension = target.rsplit(".", 1)[-1].lower()
        for node in types:
            if (node.tag.endswith("}Default")
                    and node.get("Extension", "").lower() == extension):
                return target.lstrip("/"), node.get("ContentType")
    return target.lstrip("/"), None


def package_kind(data):
    """Identify supported ZIP content independently of the filename; never raise."""
    try:
        _, content_type = package_main(data)
        if content_type in WORD_MAIN_TYPES:
            return "word"
        if content_type == "application/vnd.ms-excel.sheet.binary.macroEnabled.main":
            return "xlsb"
        if content_type == "application/vnd.ms-excel.addin.macroEnabled.main+xml":
            return "xlam"
    except Exception:
        pass
    return "excel"


def sibling_module(name):
    """Import the installed sibling build first, then the repository module."""
    import importlib

    script = Path(__file__).resolve()
    for directory in (script.with_suffix(""), script.parent):
        if (directory / (name + ".py")).is_file():
            if str(directory) not in sys.path:
                sys.path.insert(0, str(directory))
            return importlib.import_module(name)
    raise ValueError("変換モジュール " + name + ".py がありません。セットアップを再実行してください。")


def word_module():
    return sibling_module("word_ooxml")


def xlam_source(data):
    """Make openpyxl recognize an add-in without changing the original package."""
    import xml.etree.ElementTree as ET
    import zipfile

    main, _ = package_main(data)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(output, "w") as result:
        types = ET.fromstring(original.read("[Content_Types].xml"))
        for node in types:
            if (node.tag.endswith("}Override")
                    and node.get("PartName", "").lstrip("/").casefold() == main.casefold()):
                node.set("ContentType", "application/vnd.ms-excel.sheet.macroEnabled.main+xml")
                break
        else:
            namespace = types.tag.rsplit("}", 1)[0] + "}" if types.tag.startswith("{") else ""
            ET.SubElement(types, namespace + "Override", PartName="/" + main,
                          ContentType="application/vnd.ms-excel.sheet.macroEnabled.main+xml")
        for info in original.infolist():
            result.writestr(info, ET.tostring(types, encoding="utf-8")
                            if info.filename == "[Content_Types].xml" else original.read(info))
    output.seek(0)
    return output


def convert(data):
    out = io.StringIO()
    if data.startswith(b"PK"):
        kind = package_kind(data)
        if kind == "word":
            return word_module().render_text(data)
        if kind == "xlsb":
            return sibling_module("excel_xlsb").render_text(data)
        render_ooxml(xlam_source(data) if kind == "xlam" else io.BytesIO(data), out)
    elif data[:8] == OLE_SIGNATURE:
        render_xls(data, out)
    else:
        raise ValueError("対応する Excel ファイルではありません (xlsx/xlsm/xltx/xltm/xlam/xlsb/xls/xlt)")
    return out.getvalue()


def write_atomic(path, data):
    import tempfile

    handle, temporary = tempfile.mkstemp(dir=path.parent, prefix=".tmp-")
    try:
        with os.fdopen(handle, "wb") as output:
            output.write(data)
        os.replace(temporary, path)
    except OSError:
        # Another process may be reading the same entry (Windows refuses to
        # replace an open file); its content is identical, so keep it.
        try:
            os.unlink(temporary)
        except OSError:
            pass


def read_cached(path):
    try:
        data = path.read_bytes()
    except FileNotFoundError:
        return None
    try:
        os.utime(path)  # Keep recently used entries when pruning.
    except OSError:
        pass
    return data.decode("utf-8")


def prune(cache_dir):
    marker = cache_dir / "last-prune"
    now = time.time()
    try:
        if now - marker.stat().st_mtime < PRUNE_INTERVAL:
            return
    except FileNotFoundError:
        pass
    marker.touch()
    results = []
    for path in cache_dir.iterdir():
        try:
            stat = path.stat()
        except FileNotFoundError:
            continue
        age = now - stat.st_mtime
        temporary = path.name.startswith(".tmp-") and age > PRUNE_INTERVAL
        old_request = path.suffix == ".in" and age > REQUEST_MAX_AGE
        old_result = path.suffix in (".txt", ".failed") and age > CACHE_MAX_AGE
        if temporary or old_request or old_result:
            path.unlink(missing_ok=True)
        elif path.suffix == ".txt":
            results.append((stat.st_mtime, stat.st_size, path))
    total = sum(size for _, size, _ in results)
    for _, size, path in sorted(results):
        if total <= CACHE_LIMIT_BYTES:
            break
        path.unlink(missing_ok=True)
        total -= size


def timeout_seconds():
    try:
        return float(os.environ.get(TIMEOUT_VARIABLE, DEFAULT_TIMEOUT))
    except ValueError:
        return DEFAULT_TIMEOUT


def input_size(data):
    """Bytes to parse: unpacked worksheet/string parts and Word XML parts."""
    if data.startswith(b"PK"):
        import zipfile

        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if package_kind(data) == "xlsb":
                    return sum(info.file_size for info in archive.infolist()
                               if (info.filename.startswith("xl/worksheets/")
                                   and info.filename.endswith(".bin"))
                               or info.filename == "xl/sharedStrings.bin")
                return sum(
                    info.file_size for info in archive.infolist()
                    if info.filename.startswith("xl/worksheets/")
                    or info.filename == "xl/sharedStrings.xml"
                    or (info.filename.startswith("word/") and info.filename.endswith(".xml"))
                )
        except zipfile.BadZipFile:
            return 0  # convert() reports the broken file.
    return len(data)


def sheet_titles(data):
    """Read only the sheet names, which is fast even for a large workbook."""
    try:
        if data.startswith(b"PK"):
            if package_kind(data) == "xlsb":
                return sibling_module("excel_xlsb").sheet_titles(data)
            import xml.etree.ElementTree as ET
            import zipfile

            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                root = ET.fromstring(archive.read("xl/workbook.xml"))
            return [node.get("name") for node in root.iter() if node.tag.endswith("}sheet")]
        if data[:8] == OLE_SIGNATURE:
            import xlrd

            book = xlrd.open_workbook(file_contents=data, on_demand=True)
            try:
                return book.sheet_names()
            finally:
                book.release_resources()
    except Exception:
        pass
    return []


def placeholder(key, data):
    if package_kind(data) == "word":
        lines = [
            "# Word を裏で変換中 (" + key[:12] + ")",
            "# 完了後に選び直すか R で表示",
        ]
        try:
            lines += word_module().placeholder_lines(data)
        except Exception:
            pass
        return "\n".join(lines) + "\n"
    # The hash differs between the two sides of a diff, so Git still reports
    # the file as changed while the cells are being converted.
    # Short lines, so they fit one half of delta's side-by-side view.
    lines = [
        "# Excel を裏で変換中 (" + key[:12] + ")",
        "# 完了後に選び直すか R で表示",
    ]
    # The converted text starts with the same sheet line. When the other side
    # is already converted, this keeps the note at the top of the diff instead
    # of after every deleted cell.
    lines += ["Sheet: " + quoted(title) for title in sheet_titles(data)]
    return "\n".join(lines) + "\n"


def take_lock(lock):
    try:
        os.close(os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY))
    except FileExistsError:
        return False
    return True


def worker_running(lock):
    try:
        return time.time() - lock.stat().st_mtime < WORKER_STALE
    except FileNotFoundError:
        return False


def start_worker(cache_dir):
    import subprocess

    lock = cache_dir / "worker.lock"
    if worker_running(lock):
        return
    lock.unlink(missing_ok=True)
    # Take the lock for the worker before it starts, so requests arriving
    # while it is still starting up do not start more workers.
    if not take_lock(lock):
        return
    options = {}
    if os.name == "nt":
        options["creationflags"] = (
            subprocess.DETACHED_PROCESS
            | subprocess.CREATE_NEW_PROCESS_GROUP
            | subprocess.BELOW_NORMAL_PRIORITY_CLASS
        )
    else:
        options["start_new_session"] = True
    try:
        # Do not inherit Git's stdout pipe: Git would otherwise wait for the worker.
        subprocess.Popen(
            [sys.executable, str(Path(__file__).resolve()), "--worker", "--cache-dir", str(cache_dir)],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            close_fds=True,
            **options,
        )
    except OSError:
        lock.unlink(missing_ok=True)
        raise


def request_conversion(cache_dir, key, data):
    pending = cache_dir / (key + ".in")
    if pending.exists():
        os.utime(pending)  # The worker converts the latest request first.
    else:
        write_atomic(pending, data)
    start_worker(cache_dir)


def cached_convert(data, cache_dir):
    """Return (text, finished). Unfinished conversions continue in the worker."""
    key = hashlib.sha256(data).hexdigest()
    cache_dir.mkdir(parents=True, exist_ok=True)
    result = cache_dir / (key + ".txt")
    text = read_cached(result)
    if text is not None:
        return text, True
    failed = cache_dir / (key + ".failed")
    if failed.exists():
        raise ValueError(failed.read_text(encoding="utf-8"))
    timeout = timeout_seconds()
    pending = cache_dir / (key + ".in")
    if timeout > 0 and pending.exists() and worker_running(cache_dir / "worker.lock"):
        # Already queued: answer at once instead of converting it twice.
        os.utime(pending)
        return placeholder(key, data), False
    rate = (XLSB_OPTIMISTIC_BYTES_PER_SECOND if package_kind(data) == "xlsb"
            else OPTIMISTIC_BYTES_PER_SECOND)
    if timeout > 0 and input_size(data) > timeout * rate:
        # Too large to finish in time: do not spend the wait on it.
        request_conversion(cache_dir, key, data)
        return placeholder(key, data), False
    import threading

    outcome = {}

    def work():
        try:
            outcome["text"] = convert(data)
        except Exception as error:
            outcome["error"] = error

    thread = threading.Thread(target=work, daemon=True)
    thread.start()
    thread.join(timeout if timeout > 0 else None)
    if thread.is_alive():
        request_conversion(cache_dir, key, data)
        return placeholder(key, data), False
    if "error" in outcome:
        raise outcome["error"]
    write_atomic(result, outcome["text"].encode("utf-8"))
    prune(cache_dir)
    return outcome["text"], True


def next_request(cache_dir):
    newest = None
    for path in cache_dir.glob("*.in"):
        try:
            mtime = path.stat().st_mtime
        except FileNotFoundError:
            continue
        done = (cache_dir / (path.stem + ".txt")).exists()
        if done or time.time() - mtime > REQUEST_MAX_AGE:
            path.unlink(missing_ok=True)
        elif newest is None or mtime > newest[0]:
            newest = (mtime, path)
    return newest and newest[1]


def run_worker(cache_dir):
    import threading

    lock = cache_dir / "worker.lock"
    if os.name != "nt":
        os.nice(10)
    # start_worker took the lock for this process.
    lock.touch()
    owned = True
    while owned or next_request(cache_dir) is not None:
        if not owned and not take_lock(lock):
            return 0  # Another worker owns the queue.
        owned = False
        stop = threading.Event()

        def heartbeat():
            while not stop.wait(WORKER_HEARTBEAT):
                try:
                    os.utime(lock)
                except OSError:
                    pass

        threading.Thread(target=heartbeat, daemon=True).start()
        try:
            while True:
                pending = next_request(cache_dir)
                if pending is None:
                    break
                try:
                    data = pending.read_bytes()
                    write_atomic(cache_dir / (pending.stem + ".txt"), convert(data).encode("utf-8"))
                except FileNotFoundError:
                    continue
                except Exception as error:
                    write_atomic(cache_dir / (pending.stem + ".failed"), str(error).encode("utf-8"))
                finally:
                    pending.unlink(missing_ok=True)
                prune(cache_dir)
        finally:
            stop.set()
            lock.unlink(missing_ok=True)
        # Loop again: a request may have arrived after the last scan.
    return 0


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path, nargs="?", help="Excel/Word file or Git temporary blob")
    parser.add_argument("--cache-dir", type=Path,
                        help="変換結果のキャッシュ。指定すると時間のかかる変換を裏で続ける")
    parser.add_argument("--worker", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.worker:
        return run_worker(args.cache_dir) if args.cache_dir else 2
    if args.file is None:
        parser.error("Excel か Word のファイルを指定してください")
    # Git consumes UTF-8 even on Windows with a non-UTF-8 console code page.
    sys.stdout.reconfigure(encoding="utf-8", newline="\n")
    sys.stderr.reconfigure(encoding="utf-8")
    data = None
    try:
        data = args.file.read_bytes()
        if args.cache_dir is None:
            text, finished = convert(data), True
        else:
            text, finished = cached_convert(data, args.cache_dir)
    except (Exception, KeyboardInterrupt) as error:
        label = "Word" if (args.file.suffix.lower() in WORD_SUFFIXES
                           or (data is not None and package_kind(data) == "word")) else "Excel"
        text = str(error)
        if label == "Word" and data is not None:
            if data[:8] == OLE_SIGNATURE:
                text = "暗号化された Word 文書と .doc 形式は表示できません"
            elif not data.startswith(b"PK"):
                text = "対応する Word ファイルではありません (docx/docm/dotx/dotm)"
        print(label + " の差分変換に失敗しました: " + text, file=sys.stderr)
        return 1
    sys.stdout.write(text)
    sys.stdout.flush()
    if not finished:
        # Leave the abandoned in-process conversion behind without waiting.
        os._exit(0)
    return 0


if __name__ == "__main__":
    sys.exit(main())
